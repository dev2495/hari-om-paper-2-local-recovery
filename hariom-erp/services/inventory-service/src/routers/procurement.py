"""Procurement V2 receipts, commercial review, plans, alerts and RM cost control."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal
import csv
import io
import json
import uuid
from io import BytesIO
from typing import Any, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, ConfigDict, Field, field_validator
from sqlalchemy import func, select, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import (
    CostSource,
    ItemMaster,
    InventoryLocation,
    LabelPrintJob,
    LotLabelRecord,
    MrpRun,
    PaperReel,
    PlanConversion,
    ProcurementPlan,
    ProcurementPlanEntry,
    PurchaseDebitNote,
    PurchaseDebitNoteLine,
    PurchaseDebitNoteSettlement,
    PurchaseDeliverySchedule,
    PurchaseDiscrepancy,
    PurchaseOrder,
    PurchaseOrderLine,
    PurchaseOrderRevision,
    PurchaseOrderRevisionLine,
    PurchaseReceipt,
    PurchaseReceiptLine,
    ReceiptInvoiceAllocation,
    ReceiptStockAllocation,
    ReceiptScheduleAllocation,
    ReferenceType,
    ReelScanEvent,
    ReelScanEventType,
    ReelScanSource,
    ReelStatus,
    RmCostComponent,
    RmCostSheet,
    RmCostVersion,
    StockAlertEpisode,
    StockAlertPolicy,
    StockBatch,
    StockTransaction,
    SupplierInvoice,
    SupplierInvoiceLine,
    TrackingMode,
    TransactionType,
)
from ..services.labels import reel_label_payload
from ..services.receipt_quality import initial_quality, refresh_receipt_stock
from ..services.procurement import (
    ProcurementRuleError,
    alert_severity,
    calculate_landed_cost,
    canonical_hash,
    compare_rates,
    month_start,
    mrp_suggestion,
    normalize_invoice_number,
    physical_form,
)
from ..services.procurement_documents import render_debit_note_pdf
from ..services.stock_calc import get_usable_item_qty
from ..utils.auth import get_current_plant, get_current_user, require_cost_administrator, require_role
from .purchase import _actor, _next_doc_no, _next_purchase_order_no, _persist_revision


router = APIRouter(prefix="/inventory/procurement", tags=["procurement-v2"])


def _enum(value: Any) -> str:
    return str(getattr(value, "value", value))


def _plant_uuid(plant_id: str) -> uuid.UUID:
    try:
        return uuid.UUID(str(plant_id))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Procurement paper lots require a concrete UUID plant") from exc


class PaperLotInput(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)

    source_reel_no: str = Field(min_length=1, max_length=120)
    net_weight_kg: float = Field(gt=0)
    gross_weight_kg: Optional[float] = Field(default=None, gt=0)
    tare_weight_kg: Optional[float] = Field(default=None, ge=0)
    width_mm: float = Field(gt=0)
    physical_form: Optional[str] = Field(default=None, pattern="^(REEL|COIL)$")
    vendor_batch_no: Optional[str] = Field(default=None, max_length=120)
    location_id: Optional[uuid.UUID] = None


class GovernedReceiptLineInput(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)

    po_line_id: uuid.UUID
    quantity: Optional[float] = Field(default=None, gt=0)
    invoice_quantity: Optional[float] = Field(default=None, gt=0)
    invoice_rate: Optional[float] = Field(default=None, ge=0)
    tax_amount: float = Field(default=0, ge=0)
    charge_amount: float = Field(default=0, ge=0)
    batch_no: Optional[str] = Field(default=None, max_length=100)
    location_id: Optional[uuid.UUID] = None
    lots: list[PaperLotInput] = Field(default_factory=list)


class GovernedReceiptCreate(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)

    request_id: uuid.UUID
    purchase_order_id: uuid.UUID
    received_date: date
    invoice_pending: bool = False
    invoice_no: Optional[str] = Field(default=None, max_length=120)
    invoice_date: Optional[date] = None
    invoice_document_ref: Optional[str] = Field(default=None, max_length=500)
    currency: str = Field(default="INR", pattern="^[A-Z]{3}$")
    lines: list[GovernedReceiptLineInput] = Field(min_length=1)

    @field_validator("invoice_no")
    @classmethod
    def clean_invoice_no(cls, value: Optional[str]) -> Optional[str]:
        cleaned = str(value or "").strip()
        return cleaned or None


def _receipt_payload(receipt: PurchaseReceipt, db: Session) -> dict[str, Any]:
    invoice = db.query(SupplierInvoice).filter(SupplierInvoice.id == receipt.supplier_invoice_id).first() if receipt.supplier_invoice_id else None
    stock_rows = db.query(ReceiptStockAllocation).join(
        PurchaseReceiptLine, ReceiptStockAllocation.receipt_line_id == PurchaseReceiptLine.id,
    ).filter(PurchaseReceiptLine.receipt_id == receipt.id).all()
    lots = []
    for allocation in stock_rows:
        if allocation.reel_id:
            reel = db.query(PaperReel).filter(PaperReel.id == allocation.reel_id).first()
            if reel:
                label = db.query(LotLabelRecord).filter(LotLabelRecord.reel_id == reel.id).first()
                lots.append({
                    "id": str(reel.id), "at_no": reel.reel_code, "source_reel_no": reel.source_reel_no,
                    "physical_form": reel.physical_form, "net_weight_kg": float(reel.inward_weight_kg or 0),
                    "current_weight_kg": float(reel.current_weight_kg or 0), "width_mm": reel.width_mm,
                    "stock_status": reel.stock_status, "commercial_status": reel.commercial_status,
                    "label": label.content_snapshot if label else None,
                })
    return {
        "id": str(receipt.id), "request_id": str(receipt.request_id) if receipt.request_id else None,
        "purchase_order_id": str(receipt.purchase_order_id) if receipt.purchase_order_id else None,
        "receipt_kind": receipt.receipt_kind, "supplier_id": str(receipt.supplier_id or receipt.order.supplier_id) if receipt.supplier_id or receipt.order else None,
        "supplier_name": receipt.supplier_name_snapshot or (receipt.order.supplier_name_snapshot if receipt.order else None),
        "manual_reason": receipt.manual_reason, "approval_history": receipt.approval_history or [], "po_no": receipt.order.po_no if receipt.order else None,
        "grn_no": receipt.grn_no, "received_date": receipt.received_date.isoformat(), "status": receipt.status,
        "invoice_pending": receipt.invoice_pending, "commercial_status": receipt.commercial_status,
        "supplier_invoice_id": str(receipt.supplier_invoice_id) if receipt.supplier_invoice_id else None,
        "invoice_no": invoice.invoice_no if invoice else None,
        "invoice_date": invoice.invoice_date.isoformat() if invoice else None,
        "posting_version": receipt.posting_version,
        "lines": [{
            "id": str(line.id), "po_line_id": str(line.purchase_order_line_id) if line.purchase_order_line_id else None, "item_id": str(line.item_id),
            "item_code": line.item.item_code if line.item else None, "item_name": line.item.name if line.item else None,
            "quantity": float(line.qty_received or 0), "uom": line.order_line.uom if line.order_line else _enum(line.item.uom),
            "po_rate": float(line.po_rate) if line.po_rate is not None else None, "invoice_rate": float(line.invoice_rate) if line.invoice_rate is not None else None,
            "tracking_mode": line.tracking_mode, "qc_status": line.qc_status, "commercial_status": line.commercial_status,
        } for line in receipt.lines or []],
        "lots": lots, "lot_count": len(lots), "received_kg": round(sum(float(row["net_weight_kg"]) for row in lots), 3),
    }


@router.get("/receivable-lines")
def receivable_lines(
    supplier_id: Optional[uuid.UUID] = Query(default=None),
    db: Session = Depends(get_db),
    plant_id: str = Depends(get_current_plant),
    current_user: dict = Depends(get_current_user),
):
    query = db.query(PurchaseOrderLine).join(PurchaseOrder).filter(
        PurchaseOrder.plant_id == plant_id,
        PurchaseOrder.status.in_(["APPROVED", "PARTIALLY_RECEIVED"]),
        PurchaseOrderLine.qty_received + PurchaseOrderLine.qty_short_closed < PurchaseOrderLine.qty_ordered,
    )
    if supplier_id:
        query = query.filter(PurchaseOrder.supplier_id == supplier_id)
    items = []
    for line in query.order_by(PurchaseOrder.expected_date.asc().nullslast(), PurchaseOrder.po_no).all():
        items.append({
            "po_id": str(line.purchase_order_id), "po_no": line.order.po_no, "revision_no": line.order.current_revision_no,
            "supplier_id": str(line.order.supplier_id), "supplier_name": line.order.supplier_name_snapshot,
            "po_line_id": str(line.id), "logical_line_id": str(line.logical_line_id), "item_id": str(line.item_id),
            "item_code": line.item.item_code if line.item else None, "item_name": line.item.name if line.item else None,
            "tracking_mode": _enum(line.item.tracking_mode) if line.item else "BULK", "uom": line.uom,
            "ordered_qty": float(line.qty_ordered), "received_qty": float(line.qty_received),
            "open_qty": round(float(line.qty_ordered) - float(line.qty_received) - float(line.qty_short_closed), 3),
            "approved_rate": float(line.unit_cost), "expected_unit_count": line.expected_unit_count,
            "received_unit_count": line.received_unit_count, "count_basis": line.count_basis,
            "count_variance": ((line.received_unit_count or 0) - line.expected_unit_count)
                if line.received_unit_count is not None and line.expected_unit_count is not None else None,
            "specification": line.metadata_json or {},
        })
    return {"items": items}


def _validate_receipt(payload: GovernedReceiptCreate, db: Session = None, plant_id: str = None) -> None:
    if db is not None:
        location_ids = {location_id for line in payload.lines for location_id in [line.location_id, *[lot.location_id for lot in line.lots]] if location_id}
        if location_ids and db.query(InventoryLocation).filter(InventoryLocation.id.in_(location_ids), InventoryLocation.plant_id == plant_id, InventoryLocation.active == "true").count() != len(location_ids):
            raise HTTPException(422, "Choose active storage locations in the receipt plant")
    line_ids = [line.po_line_id for line in payload.lines if line.po_line_id]
    if len(line_ids) != len(set(line_ids)):
        raise HTTPException(422, "Each PO line may appear once per receipt; add its physical lots to the same line")
    if not payload.invoice_pending and any(line.invoice_rate is None for line in payload.lines):
        raise HTTPException(422, "Enter an invoice rate for every material")
    if not payload.invoice_pending and (not payload.invoice_no or not payload.invoice_date):
        raise HTTPException(status_code=422, detail="Invoice number and invoice date are required unless invoice pending is selected")
    if payload.invoice_pending and payload.invoice_no:
        raise HTTPException(status_code=422, detail="Do not enter an invoice number while invoice pending is selected")


def _width_variance(lots: list[Any], specification: dict[str, Any]) -> Optional[dict[str, Any]]:
    expected_raw = specification.get("width_mm")
    if expected_raw in {None, ""}:
        return None
    expected = Decimal(str(expected_raw))
    tolerance = Decimal(str(specification.get("width_tolerance_mm") or 0))
    violations = []
    for lot in lots:
        measured_raw = lot.get("width_mm") if isinstance(lot, dict) else getattr(lot, "width_mm", None)
        source_no = lot.get("source_reel_no") if isinstance(lot, dict) else getattr(lot, "source_reel_no", None)
        if measured_raw is None:
            continue
        measured = Decimal(str(measured_raw))
        if abs(measured - expected) > tolerance:
            violations.append({"source_reel_no": source_no, "measured_width_mm": str(measured), "delta_mm": str(measured - expected)})
    if not violations:
        return None
    return {"expected_width_mm": str(expected), "tolerance_mm": str(tolerance), "violations": violations}


@router.post("/receipts/preview")
def preview_receipt(
    payload: GovernedReceiptCreate,
    db: Session = Depends(get_db),
    plant_id: str = Depends(get_current_plant),
    current_user: dict = Depends(require_role(["Store", "PlantManager"])),
):
    _validate_receipt(payload, db, plant_id)
    order = db.query(PurchaseOrder).filter(PurchaseOrder.id == payload.purchase_order_id, PurchaseOrder.plant_id == plant_id).first()
    if not order or order.status not in {"APPROVED", "PARTIALLY_RECEIVED"}:
        raise HTTPException(status_code=409, detail="Select an approved purchase order in this plant")
    lines = {line.id: line for line in order.lines or []}
    result = []
    for index, requested in enumerate(payload.lines, start=1):
        po_line = lines.get(requested.po_line_id)
        if not po_line:
            raise HTTPException(status_code=422, detail=f"Receipt line {index} is not part of the selected PO")
        is_reel = _enum(po_line.item.tracking_mode) == "REEL"
        received = sum(Decimal(str(lot.net_weight_kg)) for lot in requested.lots) if is_reel else Decimal(str(requested.quantity or 0))
        if received <= 0 or (is_reel and not requested.lots):
            raise HTTPException(status_code=422, detail=f"Receipt line {index} needs physical lots or a positive quantity")
        open_qty = Decimal(str(po_line.qty_ordered)) - Decimal(str(po_line.qty_received)) - Decimal(str(po_line.qty_short_closed or 0))
        if received > open_qty:
            raise HTTPException(status_code=422, detail=f"Receipt line {index} exceeds the open PO balance")
        invoice_qty = received if requested.invoice_quantity is None else Decimal(str(requested.invoice_quantity))
        comparison = None if payload.invoice_pending else compare_rates(received, po_line.unit_cost, requested.invoice_rate)
        quantity_review = not payload.invoice_pending and invoice_qty != received
        result.append({
            "po_line_id": str(po_line.id), "item_code": po_line.item.item_code, "tracking_mode": _enum(po_line.item.tracking_mode),
            "uom": _enum(po_line.item.uom), "quantity_kg": float(received), "physical_unit_count": len(requested.lots), "po_rate": float(po_line.unit_cost),
            "invoice_rate": requested.invoice_rate, "rate_delta": float(comparison.delta_per_kg) if comparison else None,
            "invoice_quantity_kg": float(invoice_qty) if not payload.invoice_pending else None,
            "quantity_delta_kg": float(invoice_qty - received) if not payload.invoice_pending else None,
            "signed_difference": float(comparison.signed_difference) if comparison else None,
            "claimable_amount": float(comparison.claimable_amount) if comparison else 0,
            "commercial_status": "INVOICE_PENDING" if payload.invoice_pending else ("RATE_REVIEW" if comparison.requires_review or quantity_review else "CLEAR"),
        })
    return {"purchase_order_id": str(order.id), "po_no": order.po_no, "revision_no": order.current_revision_no,
            "invoice_pending": payload.invoice_pending, "lines": result,
            "totals": {"quantity_kg": round(sum(row["quantity_kg"] for row in result), 3),
                       "physical_unit_count": sum(row["physical_unit_count"] for row in result),
                       "claimable_amount": round(sum(row["claimable_amount"] for row in result), 2)}}


@router.post("/receipts")
def post_governed_receipt(
    payload: GovernedReceiptCreate,
    db: Session = Depends(get_db),
    plant_id: str = Depends(get_current_plant),
    current_user: dict = Depends(require_role(["Store", "PlantManager"])),
):
    _validate_receipt(payload, db, plant_id)
    fingerprint = canonical_hash(payload.model_dump(mode="json", exclude={"request_id"}))
    existing = db.query(PurchaseReceipt).filter(PurchaseReceipt.plant_id == plant_id, PurchaseReceipt.request_id == payload.request_id).first()
    if existing:
        if existing.request_fingerprint != fingerprint:
            raise HTTPException(status_code=409, detail="This receipt request key already exists with different details")
        response = _receipt_payload(existing, db)
        response["idempotent"] = True
        return response
    order = db.query(PurchaseOrder).filter(PurchaseOrder.id == payload.purchase_order_id, PurchaseOrder.plant_id == plant_id).with_for_update().first()
    if not order or order.status not in {"APPROVED", "PARTIALLY_RECEIVED"}:
        raise HTTPException(status_code=409, detail="Only an approved purchase order can receive material")
    revision = db.query(PurchaseOrderRevision).filter(
        PurchaseOrderRevision.purchase_order_id == order.id,
        PurchaseOrderRevision.revision_no == order.current_revision_no,
        PurchaseOrderRevision.approval_state == "APPROVED",
    ).first()
    if not revision:
        raise HTTPException(status_code=409, detail="The current PO revision is not approved")
    invoice = None
    if not payload.invoice_pending:
        normalized = normalize_invoice_number(payload.invoice_no or "")
        invoice = db.query(SupplierInvoice).filter(
            SupplierInvoice.plant_id == plant_id, SupplierInvoice.supplier_id == order.supplier_id,
            SupplierInvoice.normalized_invoice_no == normalized,
        ).first()
        if invoice and (invoice.invoice_date != payload.invoice_date or invoice.currency != payload.currency):
            raise HTTPException(status_code=409, detail="Invoice number already exists with a different date or currency")
        if not invoice:
            invoice = SupplierInvoice(
                plant_id=plant_id, supplier_id=order.supplier_id, invoice_no=payload.invoice_no or "",
                normalized_invoice_no=normalized, invoice_date=payload.invoice_date, currency=payload.currency,
                document_ref=payload.invoice_document_ref, created_by=_actor(current_user),
            )
            db.add(invoice)
            db.flush()
    grn_no = _next_doc_no(db, PurchaseReceipt, plant_id, "grn_no", "GRN")
    receipt = PurchaseReceipt(
        plant_id=plant_id, purchase_order_id=order.id, supplier_id=order.supplier_id,
        supplier_name_snapshot=order.supplier_name_snapshot, request_id=payload.request_id,
        request_fingerprint=fingerprint, supplier_invoice_id=invoice.id if invoice else None,
        invoice_pending=payload.invoice_pending, commercial_status="INVOICE_PENDING" if payload.invoice_pending else "CLEAR",
        grn_no=grn_no, received_date=payload.received_date, status="POSTED", created_by=_actor(current_user),
    )
    db.add(receipt)
    db.flush()
    order_lines = {line.id: line for line in order.lines or []}
    revision_lines = {line.source_order_line_id: line for line in revision.lines or []}
    result_lots: list[dict[str, Any]] = []
    receipt_statuses: list[str] = []
    commercial_states: list[str] = []
    for index, requested in enumerate(payload.lines, start=1):
        po_line = order_lines.get(requested.po_line_id)
        revision_line = revision_lines.get(requested.po_line_id)
        if not po_line or not revision_line:
            raise HTTPException(status_code=422, detail=f"Receipt line {index} is not in the approved revision")
        is_reel = _enum(po_line.item.tracking_mode) == "REEL"
        qty = sum(Decimal(str(lot.net_weight_kg)) for lot in requested.lots) if is_reel else Decimal(str(requested.quantity or 0))
        if qty <= 0 or (is_reel and not requested.lots) or (not is_reel and requested.lots):
            raise HTTPException(status_code=422, detail=f"Receipt line {index} does not match the item's tracking mode")
        if _enum(po_line.item.type) == "RAW_PAPER" and po_line.uom != "KG":
            raise HTTPException(status_code=409, detail=f"Approved paper line {index} is not in KG")
        remaining = Decimal(str(po_line.qty_ordered)) - Decimal(str(po_line.qty_received)) - Decimal(str(po_line.qty_short_closed or 0))
        if qty > remaining:
            raise HTTPException(status_code=422, detail=f"Receipt line {index} exceeds open PO kg")
        invoice_qty = qty if requested.invoice_quantity is None else Decimal(str(requested.invoice_quantity))
        comparison = None if payload.invoice_pending else compare_rates(qty, revision_line.unit_rate, requested.invoice_rate)
        width_evidence = _width_variance(requested.lots, revision_line.specification_json or {}) if is_reel else None
        commercial = "INVOICE_PENDING" if payload.invoice_pending else ("RATE_REVIEW" if comparison.requires_review or invoice_qty != qty or width_evidence else "CLEAR")
        quality_metadata, qc_status = initial_quality(po_line.item, plant_id, payload.received_date)
        stock_status = "BLOCKED" if commercial != "CLEAR" else ("QC_HOLD" if qc_status == "PENDING" else "UNRESTRICTED")
        receipt_line = PurchaseReceiptLine(
            receipt_id=receipt.id, purchase_order_line_id=po_line.id, approved_revision_line_id=revision_line.id,
            item_id=po_line.item_id, qty_received=qty, unit_cost=float(revision_line.unit_rate),
            po_rate=revision_line.unit_rate, invoice_rate=requested.invoice_rate,
            tracking_mode="REEL" if is_reel else "BULK", commercial_status=commercial, qc_status=qc_status,
        )
        db.add(receipt_line)
        db.flush()
        if is_reel:
            seen_source: set[str] = set()
            for lot_index, lot in enumerate(requested.lots, start=1):
                source_no = lot.source_reel_no.strip().upper()
                if source_no in seen_source:
                    raise HTTPException(status_code=422, detail=f"Duplicate vendor reel number in receipt line {index}: {source_no}")
                seen_source.add(source_no)
                duplicate = db.query(PaperReel).filter(
                    PaperReel.plant_id == _plant_uuid(plant_id), PaperReel.supplier_id == order.supplier_id,
                    PaperReel.source_reel_no == source_no,
                ).first()
                if duplicate:
                    raise HTTPException(status_code=409, detail=f"Vendor reel {source_no} is already inwarded as {duplicate.reel_code}")
                if lot.gross_weight_kg is not None and lot.tare_weight_kg is not None:
                    measured_net = round(lot.gross_weight_kg - lot.tare_weight_kg, 3)
                    if abs(measured_net - lot.net_weight_kg) > 0.001:
                        raise HTTPException(status_code=422, detail=f"Lot {lot_index} gross minus tare does not equal net kg")
                form = lot.physical_form or physical_form(lot.width_mm)
                at_no = _next_doc_no(db, PaperReel, _plant_uuid(plant_id), "reel_code", "AT")
                reel = PaperReel(
                    plant_id=_plant_uuid(plant_id), reel_code=at_no, purchase_receipt_line_id=receipt_line.id,
                    paper_id=po_line.item_id, gsm=(po_line.metadata_json or {}).get("gsm"),
                    bf=(po_line.metadata_json or {}).get("plybond"), supplier_id=order.supplier_id,
                    supplier_name=order.supplier_name_snapshot, supplier_name_snapshot=order.supplier_name_snapshot,
                    inward_weight_kg=lot.net_weight_kg, net_weight_kg=lot.net_weight_kg,
                    gross_weight_kg=lot.gross_weight_kg, tare_weight_kg=lot.tare_weight_kg,
                    current_weight_kg=lot.net_weight_kg, physical_form=form, source_reel_no=source_no,
                    vendor_batch_no=lot.vendor_batch_no, width_mm=lot.width_mm,
                    unit_cost=float(revision_line.unit_rate),
                    cost_source=CostSource.SUPPLIER, status=ReelStatus.IN_STOCK, stock_status=stock_status,
                    commercial_status=commercial, location_id=lot.location_id, inward_date=payload.received_date,
                    inward_metadata={**quality_metadata, "po_id": str(order.id), "po_no": order.po_no, "po_revision": revision.revision_no,
                        "po_line_id": str(po_line.id), "receipt_id": str(receipt.id), "grn_no": receipt.grn_no,
                        "invoice_no": payload.invoice_no, "invoice_date": payload.invoice_date.isoformat() if payload.invoice_date else None,
                        "po_rate": str(revision_line.unit_rate), "invoice_rate": requested.invoice_rate,
                        "source_reel_no": source_no, "physical_form": form},
                )
                db.add(reel)
                db.flush()
                db.add(ReceiptStockAllocation(receipt_line_id=receipt_line.id, reel_id=reel.id, allocated_qty=lot.net_weight_kg))
                label_snapshot = reel_label_payload(reel, po_line.item)
                db.add(LotLabelRecord(plant_id=plant_id, reel_id=reel.id, label_code=at_no, content_snapshot=label_snapshot))
                db.add(ReelScanEvent(plant_id=_plant_uuid(plant_id), reel_id=reel.id, event_type=ReelScanEventType.INWARD_SCAN,
                    source=ReelScanSource.INVENTORY, event_metadata={"receipt_id": str(receipt.id), "at_no": at_no, "source_reel_no": source_no}))
                result_lots.append({"id": str(reel.id), "at_no": at_no, "source_reel_no": source_no,
                    "physical_form": form, "width_mm": lot.width_mm, "net_weight_kg": lot.net_weight_kg, "stock_status": stock_status,
                    "label": label_snapshot})
        else:
            batch_no = (requested.batch_no or f"{grn_no}-B{index:03d}").strip().upper()
            batch = StockBatch(item_id=po_line.item_id, batch_no=batch_no, received_qty=float(qty),
                location_id=requested.location_id, stock_status=stock_status,
                unit_cost=float(revision_line.unit_rate),
                cost_source="SUPPLIER", supplier_id=order.supplier_id, supplier_name_snapshot=order.supplier_name_snapshot,
                plant_id=plant_id, inward_metadata={**quality_metadata, "po_id": str(order.id), "receipt_id": str(receipt.id), "invoice_no": payload.invoice_no})
            db.add(batch)
            db.flush()
            receipt_line.batch_id = batch.id
            db.add(ReceiptStockAllocation(receipt_line_id=receipt_line.id, batch_id=batch.id, allocated_qty=qty))
            db.add(StockTransaction(item_id=po_line.item_id, batch_id=batch.id, transaction_type=TransactionType.INWARD,
                qty_change=float(qty), reference_type=ReferenceType.PURCHASE, reference_id=order.id, plant_id=plant_id,
                effective_date=payload.received_date, location_id=requested.location_id, stock_status=stock_status,
                external_ref=f"PGRN:{receipt.id}:{index}", movement_metadata={"receipt_id": str(receipt.id), "po_revision": revision.revision_no}))
        if invoice:
            invoice_line = SupplierInvoiceLine(invoice_id=invoice.id, item_id=po_line.item_id, qty=invoice_qty,
                rate=requested.invoice_rate, uom=po_line.uom, tax_amount=requested.tax_amount, charge_amount=requested.charge_amount)
            db.add(invoice_line)
            db.flush()
            allocation = ReceiptInvoiceAllocation(receipt_line_id=receipt_line.id, invoice_line_id=invoice_line.id,
                revision_line_id=revision_line.id, allocated_qty=qty, po_rate=revision_line.unit_rate,
                invoice_rate=requested.invoice_rate, rate_delta=comparison.delta_per_kg,
                claimable_amount=comparison.claimable_amount)
            db.add(allocation)
            db.flush()
            if comparison.requires_review:
                db.add(PurchaseDiscrepancy(plant_id=plant_id, allocation_id=allocation.id, discrepancy_type="RATE",
                    quantity=qty, po_rate=revision_line.unit_rate, invoice_rate=requested.invoice_rate,
                    delta=comparison.delta_per_kg, claimable_amount=comparison.claimable_amount, status="OPEN"))
            if invoice_qty != qty:
                db.add(PurchaseDiscrepancy(plant_id=plant_id, allocation_id=allocation.id, discrepancy_type="QUANTITY",
                    quantity=qty, po_rate=revision_line.unit_rate, invoice_rate=requested.invoice_rate,
                    delta=invoice_qty - qty, claimable_amount=0, status="OPEN"))
            if width_evidence:
                first = width_evidence["violations"][0]
                db.add(PurchaseDiscrepancy(plant_id=plant_id, allocation_id=allocation.id, discrepancy_type="SPECIFICATION",
                    quantity=qty, po_rate=width_evidence["expected_width_mm"], invoice_rate=first["measured_width_mm"],
                    delta=first["delta_mm"], claimable_amount=0, evidence_json=width_evidence, status="OPEN"))
        po_line.qty_received = float(po_line.qty_received or 0) + float(qty)
        if is_reel:
            po_line.received_unit_count = int(po_line.received_unit_count or 0) + len(requested.lots)
        po_line.line_status = "CLOSED" if po_line.qty_received + po_line.qty_short_closed + 1e-9 >= po_line.qty_ordered else "PARTIAL"
        schedule_remaining = qty
        schedules = db.query(PurchaseDeliverySchedule).filter(
            PurchaseDeliverySchedule.purchase_order_line_id == po_line.id,
            PurchaseDeliverySchedule.confirmation_status != "CANCELLED",
        ).order_by(PurchaseDeliverySchedule.delivery_date, PurchaseDeliverySchedule.created_at).with_for_update().all()
        for schedule in schedules:
            schedule_open = Decimal(str(schedule.planned_qty)) - Decimal(str(schedule.received_qty or 0)) - Decimal(str(schedule.cancelled_qty or 0))
            if schedule_open <= 0 or schedule_remaining <= 0:
                continue
            applied = min(schedule_open, schedule_remaining)
            db.add(ReceiptScheduleAllocation(plant_id=plant_id, receipt_line_id=receipt_line.id,
                schedule_id=schedule.id, allocated_qty=float(applied)))
            schedule.version += 1
            schedule_remaining -= applied
        receipt_statuses.append(qc_status)
        commercial_states.append(commercial)
    order.status = "RECEIVED" if all(line.line_status == "CLOSED" for line in order.lines or []) else "PARTIALLY_RECEIVED"
    receipt.status = "QC_PENDING" if "PENDING" in receipt_statuses else "POSTED"
    receipt.commercial_status = (
        "INVOICE_PENDING" if "INVOICE_PENDING" in commercial_states
        else "REVIEW_REQUIRED" if "RATE_REVIEW" in commercial_states
        else "CLEAR"
    )
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Receipt could not be posted because an invoice, lot, batch or request reference already exists") from exc
    db.refresh(receipt)
    response = _receipt_payload(receipt, db)
    response["created_lots"] = result_lots
    return response


@router.get("/receipts")
def list_governed_receipts(
    limit: int = Query(default=100, ge=1, le=500),
    offset: int = Query(default=0, ge=0), q: Optional[str] = Query(default=None, max_length=200),
    receipt_kind: Optional[str] = Query(default=None, pattern="^(MANUAL|PO_LINKED)$"),
    invoice_pending: Optional[bool] = Query(default=None),
    db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
    current_user: dict = Depends(get_current_user),
):
    query = db.query(PurchaseReceipt).outerjoin(PurchaseOrder, PurchaseOrder.id == PurchaseReceipt.purchase_order_id).outerjoin(SupplierInvoice, SupplierInvoice.id == PurchaseReceipt.supplier_invoice_id).filter(PurchaseReceipt.plant_id == plant_id)
    if receipt_kind:
        query = query.filter(PurchaseReceipt.receipt_kind == receipt_kind)
    if invoice_pending is not None:
        query = query.filter(PurchaseReceipt.invoice_pending == invoice_pending)
    if q and q.strip():
        term = "%" + q.strip() + "%"
        query = query.filter(or_(PurchaseReceipt.grn_no.ilike(term), PurchaseReceipt.supplier_name_snapshot.ilike(term), PurchaseOrder.po_no.ilike(term), SupplierInvoice.invoice_no.ilike(term)))
    total = query.count()
    rows = query.order_by(PurchaseReceipt.created_at.desc(), PurchaseReceipt.id).offset(offset).limit(limit).all()
    return {"items": [_receipt_payload(row, db) for row in rows], "total": total, "offset": offset, "limit": limit}


@router.get("/receipts/{receipt_id}")
def get_governed_receipt(receipt_id: uuid.UUID, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    receipt = db.query(PurchaseReceipt).filter_by(id=receipt_id, plant_id=plant_id).first()
    if not receipt:
        raise HTTPException(404, "Receipt not found in this plant")
    return _receipt_payload(receipt, db)


class AttachedInvoiceLine(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)

    receipt_line_id: uuid.UUID
    invoice_rate: float = Field(ge=0)
    invoice_quantity: Optional[float] = Field(default=None, gt=0)
    tax_amount: float = Field(default=0, ge=0)
    charge_amount: float = Field(default=0, ge=0)


class AttachInvoiceRequest(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)

    expected_version: int = Field(ge=1)
    invoice_no: str = Field(min_length=1, max_length=120)
    invoice_date: date
    currency: str = Field(default="INR", pattern="^[A-Z]{3}$")
    invoice_document_ref: Optional[str] = Field(default=None, max_length=500)
    lines: list[AttachedInvoiceLine] = Field(min_length=1)

    @field_validator("invoice_no")
    @classmethod
    def clean_invoice_no(cls, value: str) -> str:
        cleaned = value.strip()
        if not cleaned:
            raise ValueError("invoice number is required")
        return cleaned


@router.post("/receipts/{receipt_id}/attach-invoice")
def attach_invoice_to_receipt(
    receipt_id: uuid.UUID,
    payload: AttachInvoiceRequest,
    db: Session = Depends(get_db),
    plant_id: str = Depends(get_current_plant),
    current_user: dict = Depends(require_role(["Store", "PlantManager"])),
):
    """Attach a delayed invoice and run the same approved-revision comparison.

    Every saved receipt line must be represented exactly once. The physical
    receipt remains immutable; this step creates commercial allocations and
    updates only the commercial/eligibility projection of its stock objects.
    """
    receipt = db.query(PurchaseReceipt).filter(
        PurchaseReceipt.id == receipt_id,
        PurchaseReceipt.plant_id == plant_id,
    ).with_for_update().first()
    if not receipt:
        raise HTTPException(status_code=404, detail="Receipt not found")

    normalized = normalize_invoice_number(payload.invoice_no)
    if not receipt.invoice_pending:
        existing_invoice = db.query(SupplierInvoice).filter(SupplierInvoice.id == receipt.supplier_invoice_id).first()
        saved_lines = {line.id: line for line in receipt.lines or []}
        exact = bool(existing_invoice and existing_invoice.normalized_invoice_no == normalized
                     and existing_invoice.invoice_date == payload.invoice_date and existing_invoice.currency == payload.currency
                     and existing_invoice.document_ref == payload.invoice_document_ref
                     and len(payload.lines) == len(saved_lines) and {line.receipt_line_id for line in payload.lines} == set(saved_lines))
        for requested in payload.lines:
            if not exact:
                break
            saved = saved_lines[requested.receipt_line_id]
            pair = db.query(ReceiptInvoiceAllocation, SupplierInvoiceLine).join(
                SupplierInvoiceLine, SupplierInvoiceLine.id == ReceiptInvoiceAllocation.invoice_line_id,
            ).filter(ReceiptInvoiceAllocation.receipt_line_id == saved.id, SupplierInvoiceLine.invoice_id == existing_invoice.id).first()
            invoice_qty = requested.invoice_quantity if requested.invoice_quantity is not None else saved.qty_received
            exact = bool(pair and Decimal(str(pair[1].rate)) == Decimal(str(requested.invoice_rate))
                         and Decimal(str(pair[1].qty)) == Decimal(str(invoice_qty))
                         and Decimal(str(pair[1].tax_amount)) == Decimal(str(requested.tax_amount))
                         and Decimal(str(pair[1].charge_amount)) == Decimal(str(requested.charge_amount)))
        if exact:
            return {**_receipt_payload(receipt, db), "idempotent": True}
        raise HTTPException(status_code=409, detail="This receipt already has an invoice with different details; refresh and use commercial review")
    if receipt.posting_version != payload.expected_version:
        raise HTTPException(status_code=409, detail="Receipt changed; refresh before attaching the invoice")

    receipt_lines = {line.id: line for line in receipt.lines or []}
    supplied = {line.receipt_line_id: line for line in payload.lines}
    if len(supplied) != len(payload.lines):
        raise HTTPException(status_code=422, detail="Each receipt line may appear only once")
    if set(supplied) != set(receipt_lines):
        raise HTTPException(status_code=422, detail="Invoice must cover every saved receipt line exactly once")

    order = receipt.order
    supplier_id = receipt.supplier_id or order.supplier_id
    invoice = db.query(SupplierInvoice).filter(
        SupplierInvoice.plant_id == plant_id,
        SupplierInvoice.supplier_id == supplier_id,
        SupplierInvoice.normalized_invoice_no == normalized,
    ).first()
    if invoice:
        if invoice.invoice_date != payload.invoice_date or invoice.currency != payload.currency:
            raise HTTPException(status_code=409, detail="Invoice number already exists with a different date or currency")
    else:
        invoice = SupplierInvoice(
            plant_id=plant_id,
            supplier_id=supplier_id,
            invoice_no=payload.invoice_no,
            normalized_invoice_no=normalized,
            invoice_date=payload.invoice_date,
            currency=payload.currency,
            document_ref=payload.invoice_document_ref,
            created_by=_actor(current_user),
        )
        db.add(invoice)
        db.flush()

    review_required = False
    for receipt_line_id, requested in supplied.items():
        receipt_line = receipt_lines[receipt_line_id]
        qty = Decimal(str(receipt_line.qty_received))
        invoice_qty = Decimal(str(requested.invoice_quantity if requested.invoice_quantity is not None else receipt_line.qty_received))
        comparison = compare_rates(qty, receipt_line.po_rate if receipt_line.po_rate is not None else requested.invoice_rate, requested.invoice_rate)
        lot_widths = []
        for stock_allocation in db.query(ReceiptStockAllocation).filter(ReceiptStockAllocation.receipt_line_id == receipt_line.id).all():
            if stock_allocation.reel_id:
                saved_reel = db.query(PaperReel).filter(PaperReel.id == stock_allocation.reel_id).first()
                if saved_reel:
                    lot_widths.append({"source_reel_no": saved_reel.source_reel_no, "width_mm": saved_reel.width_mm})
        revision_line = db.query(PurchaseOrderRevisionLine).filter(PurchaseOrderRevisionLine.id == receipt_line.approved_revision_line_id).first()
        width_evidence = _width_variance(lot_widths, revision_line.specification_json or {}) if revision_line else None
        line_state = "MANUAL_REVIEW" if receipt.receipt_kind == "MANUAL" else ("RATE_REVIEW" if comparison.requires_review or invoice_qty != qty or width_evidence else "CLEAR")
        review_required = review_required or comparison.requires_review or invoice_qty != qty or bool(width_evidence)

        invoice_line = SupplierInvoiceLine(
            invoice_id=invoice.id,
            item_id=receipt_line.item_id,
            qty=invoice_qty,
            rate=requested.invoice_rate,
            uom=receipt_line.order_line.uom if receipt_line.order_line else _enum(receipt_line.item.uom),
            tax_amount=requested.tax_amount,
            charge_amount=requested.charge_amount,
        )
        db.add(invoice_line)
        db.flush()
        allocation = ReceiptInvoiceAllocation(
            receipt_line_id=receipt_line.id,
            invoice_line_id=invoice_line.id,
            revision_line_id=receipt_line.approved_revision_line_id,
            allocated_qty=qty,
            po_rate=receipt_line.po_rate if receipt_line.po_rate is not None else requested.invoice_rate,
            invoice_rate=requested.invoice_rate,
            rate_delta=comparison.delta_per_kg,
            claimable_amount=comparison.claimable_amount,
        )
        db.add(allocation)
        db.flush()
        if comparison.requires_review:
            db.add(PurchaseDiscrepancy(
                plant_id=plant_id,
                allocation_id=allocation.id,
                discrepancy_type="RATE",
                quantity=qty,
                po_rate=receipt_line.po_rate if receipt_line.po_rate is not None else requested.invoice_rate,
                invoice_rate=requested.invoice_rate,
                delta=comparison.delta_per_kg,
                claimable_amount=comparison.claimable_amount,
                status="OPEN",
            ))
        if invoice_qty != qty:
            db.add(PurchaseDiscrepancy(
                plant_id=plant_id,
                allocation_id=allocation.id,
                discrepancy_type="QUANTITY",
                quantity=qty,
                po_rate=receipt_line.po_rate if receipt_line.po_rate is not None else requested.invoice_rate,
                invoice_rate=requested.invoice_rate,
                delta=invoice_qty - qty,
                claimable_amount=0,
                status="OPEN",
            ))
        if width_evidence:
            first = width_evidence["violations"][0]
            db.add(PurchaseDiscrepancy(
                plant_id=plant_id, allocation_id=allocation.id, discrepancy_type="SPECIFICATION", quantity=qty,
                po_rate=width_evidence["expected_width_mm"], invoice_rate=first["measured_width_mm"],
                delta=first["delta_mm"], claimable_amount=0, evidence_json=width_evidence, status="OPEN",
            ))

        receipt_line.invoice_rate = requested.invoice_rate
        receipt_line.commercial_status = line_state
        stock_state = "BLOCKED" if line_state != "CLEAR" or invoice_qty != qty else ("QC_HOLD" if receipt_line.qc_status == "PENDING" else "UNRESTRICTED")
        for stock_allocation in db.query(ReceiptStockAllocation).filter(ReceiptStockAllocation.receipt_line_id == receipt_line.id).all():
            if stock_allocation.reel_id:
                reel = db.query(PaperReel).filter(PaperReel.id == stock_allocation.reel_id).first()
                if reel:
                    reel.unit_cost = float(receipt_line.po_rate if receipt_line.po_rate is not None else requested.invoice_rate)
                    reel.cost_source = CostSource.SUPPLIER
                    reel.commercial_status = receipt_line.commercial_status
                    reel.inward_metadata = {**(reel.inward_metadata or {}), "invoice_no": invoice.invoice_no,
                        "invoice_date": invoice.invoice_date.isoformat(), "invoice_rate": requested.invoice_rate}
            elif stock_allocation.batch_id:
                batch = db.query(StockBatch).filter(StockBatch.id == stock_allocation.batch_id).first()
                if batch:
                    batch.unit_cost = float(receipt_line.po_rate if receipt_line.po_rate is not None else requested.invoice_rate)
                    batch.cost_source = "SUPPLIER"
                    batch.inward_metadata = {**(batch.inward_metadata or {}), "invoice_no": invoice.invoice_no,
                        "invoice_date": invoice.invoice_date.isoformat(), "invoice_rate": requested.invoice_rate}

        db.flush()
        refresh_receipt_stock(db, receipt_line)

    receipt.supplier_invoice_id = invoice.id
    receipt.invoice_pending = False
    receipt.commercial_status = "MANUAL_REVIEW" if receipt.receipt_kind == "MANUAL" else ("REVIEW_REQUIRED" if review_required else "CLEAR")
    receipt.posting_version += 1
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Invoice attachment conflicts with an existing invoice or discrepancy") from exc
    db.refresh(receipt)
    return _receipt_payload(receipt, db)


def _discrepancy_payload(row: PurchaseDiscrepancy, db: Session) -> dict[str, Any]:
    allocation = db.query(ReceiptInvoiceAllocation).filter(ReceiptInvoiceAllocation.id == row.allocation_id).first()
    receipt_line = db.query(PurchaseReceiptLine).filter(PurchaseReceiptLine.id == allocation.receipt_line_id).first() if allocation else None
    receipt = db.query(PurchaseReceipt).filter(PurchaseReceipt.id == receipt_line.receipt_id).first() if receipt_line else None
    invoice_line = db.query(SupplierInvoiceLine).filter(SupplierInvoiceLine.id == allocation.invoice_line_id).first() if allocation else None
    invoice = db.query(SupplierInvoice).filter(SupplierInvoice.id == invoice_line.invoice_id).first() if invoice_line else None
    order = receipt.order if receipt else None
    item = receipt_line.item if receipt_line else None
    return {"id": str(row.id), "type": row.discrepancy_type, "status": row.status, "version": row.version,
        "po_id": str(order.id) if order else None, "po_no": order.po_no if order else None,
        "revision_no": order.current_revision_no if order else None, "receipt_id": str(receipt.id) if receipt else None,
        "grn_no": receipt.grn_no if receipt else None, "invoice_no": invoice.invoice_no if invoice else None,
        "invoice_date": invoice.invoice_date.isoformat() if invoice else None, "supplier_id": str(order.supplier_id) if order else str(receipt.supplier_id) if receipt and receipt.supplier_id else None,
        "supplier_name": order.supplier_name_snapshot if order else receipt.supplier_name_snapshot if receipt else None, "item_id": str(item.id) if item else None,
        "item_code": item.item_code if item else None, "item_name": item.name if item else None,
        "quantity_kg": float(row.quantity), "invoice_quantity_kg": float(invoice_line.qty) if invoice_line else None,
        "quantity_delta_kg": float(row.delta) if row.discrepancy_type == "QUANTITY" else 0,
        "po_rate": float(row.po_rate), "invoice_rate": float(row.invoice_rate),
        "delta": float(row.delta),
        "signed_amount": float(Decimal(row.quantity) * Decimal(row.delta)) if row.discrepancy_type == "RATE" else 0,
        "claimable_amount": float(row.claimable_amount), "assignee": row.assignee,
        "resolution_reason": row.resolution_reason, "evidence": row.evidence_json or {},
        "created_at": row.created_at.isoformat()}


@router.get("/discrepancies")
def list_discrepancies(
    status: Optional[str] = Query(default=None), limit: int = Query(default=250, ge=1, le=1000),
    db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user),
):
    query = db.query(PurchaseDiscrepancy).filter(PurchaseDiscrepancy.plant_id == plant_id)
    if status:
        query = query.filter(PurchaseDiscrepancy.status == status.strip().upper())
    rows = query.order_by(PurchaseDiscrepancy.created_at.desc()).limit(limit).all()
    return {"items": [_discrepancy_payload(row, db) for row in rows],
        "summary": {"open_cases": sum(row.status in {"OPEN", "UNDER_REVIEW"} for row in rows),
                    "claimable_amount": round(sum(float(row.claimable_amount) for row in rows if row.status not in {"RESOLVED", "REJECTED"}), 2)}}


class DiscrepancyAction(BaseModel):
    model_config = ConfigDict(extra="forbid")
    action: str = Field(pattern="^(ASSIGN|ACCEPT_VARIANCE|REJECT_INVOICE|RELEASE_STOCK)$")
    expected_version: int = Field(gt=0)
    reason: str = Field(min_length=3, max_length=1000)
    assignee: Optional[str] = Field(default=None, max_length=200)


def _refresh_receipt_line_stock(db: Session, line: PurchaseReceiptLine) -> None:
    refresh_receipt_stock(db, line)


def _refresh_receipt_commercial_status(db: Session, line: PurchaseReceiptLine) -> None:
    if line.receipt.receipt_kind == "MANUAL" and not any(event.get("action") == "MANUAL_RECEIPT_APPROVED" for event in (line.receipt.approval_history or [])):
        line.commercial_status = "MANUAL_REVIEW"
        line.receipt.commercial_status = "MANUAL_REVIEW"
        return
    allocations = db.query(ReceiptInvoiceAllocation).filter(ReceiptInvoiceAllocation.receipt_line_id == line.id).all()
    allocation_ids = [row.id for row in allocations]
    unresolved = db.query(PurchaseDiscrepancy).filter(
        PurchaseDiscrepancy.allocation_id.in_(allocation_ids),
        PurchaseDiscrepancy.status.in_(["OPEN", "UNDER_REVIEW", "CLAIM_DRAFTED", "CLAIMED"]),
    ).count() if allocation_ids else 0
    if line.commercial_status != "RELEASED_WITH_CLAIM":
        line.commercial_status = "RATE_REVIEW" if unresolved else "CLEAR"
    receipt = line.receipt
    if receipt:
        states = [candidate.commercial_status for candidate in receipt.lines or []]
        receipt.commercial_status = "REVIEW_REQUIRED" if any(state == "RATE_REVIEW" for state in states) else (
            "RELEASED_WITH_CLAIM" if any(state == "RELEASED_WITH_CLAIM" for state in states) else "CLEAR"
        )


@router.post("/discrepancies/{discrepancy_id}/action")
def act_on_discrepancy(
    discrepancy_id: uuid.UUID, payload: DiscrepancyAction, db: Session = Depends(get_db),
    plant_id: str = Depends(get_current_plant), current_user: dict = Depends(require_role(["PlantManager", "Accounts"])),
):
    row = db.query(PurchaseDiscrepancy).filter(PurchaseDiscrepancy.id == discrepancy_id, PurchaseDiscrepancy.plant_id == plant_id).with_for_update().first()
    if not row:
        raise HTTPException(status_code=404, detail="Discrepancy not found")
    if row.version != payload.expected_version:
        raise HTTPException(status_code=409, detail="Discrepancy changed; reload before deciding")
    allocation = db.query(ReceiptInvoiceAllocation).filter(ReceiptInvoiceAllocation.id == row.allocation_id).first()
    receipt_line = db.query(PurchaseReceiptLine).filter(PurchaseReceiptLine.id == allocation.receipt_line_id).first()
    if payload.action in {"ACCEPT_VARIANCE", "RELEASE_STOCK"} and receipt_line.receipt.created_by == _actor(current_user):
        raise HTTPException(status_code=409, detail="The receipt maker cannot approve or release their own commercial exception")
    if payload.action == "ASSIGN":
        if not payload.assignee:
            raise HTTPException(status_code=422, detail="Assignee is required")
        row.assignee = payload.assignee.strip()
        row.status = "UNDER_REVIEW"
    elif payload.action == "ACCEPT_VARIANCE":
        if row.status not in {"OPEN", "UNDER_REVIEW"}:
            raise HTTPException(status_code=409, detail="Only an open commercial exception can be accepted")
        if row.discrepancy_type == "RATE":
            receipt_line.unit_cost = float(row.invoice_rate)
            for stock in db.query(ReceiptStockAllocation).filter_by(receipt_line_id=receipt_line.id).all():
                target = db.get(PaperReel, stock.reel_id) if stock.reel_id else db.get(StockBatch, stock.batch_id)
                if target:
                    target.unit_cost = float(row.invoice_rate)
        row.status = "ACCEPTED"
        row.resolution_reason = payload.reason
        row.resolved_at = datetime.utcnow()
        db.flush()
        _refresh_receipt_commercial_status(db, receipt_line)
        _refresh_receipt_line_stock(db, receipt_line)
    elif payload.action == "REJECT_INVOICE":
        row.status = "UNDER_REVIEW"
        row.resolution_reason = payload.reason
    elif payload.action == "RELEASE_STOCK":
        if receipt_line.receipt.receipt_kind == "MANUAL" and not any(event.get("action") == "MANUAL_RECEIPT_APPROVED" for event in (receipt_line.receipt.approval_history or [])):
            raise HTTPException(409, "Approve the manual purchase before releasing its stock")
        if row.status not in {"ACCEPTED", "CLAIMED", "RESOLVED"}:
            raise HTTPException(status_code=409, detail="Accept the variance or issue the approved claim before releasing stock")
        sibling_allocations = select(ReceiptInvoiceAllocation.id).where(
            ReceiptInvoiceAllocation.receipt_line_id == receipt_line.id,
        )
        unresolved_siblings = db.query(PurchaseDiscrepancy).filter(
            PurchaseDiscrepancy.allocation_id.in_(sibling_allocations),
            PurchaseDiscrepancy.id != row.id,
            PurchaseDiscrepancy.status.in_(["OPEN", "UNDER_REVIEW", "CLAIM_DRAFTED"]),
        ).count()
        if unresolved_siblings:
            raise HTTPException(status_code=409, detail="Resolve every other discrepancy on this receipt line before releasing stock")
        row.resolution_reason = payload.reason
        receipt_line.commercial_status = "RELEASED_WITH_CLAIM"
        if receipt_line.receipt:
            receipt_line.receipt.commercial_status = "RELEASED_WITH_CLAIM"
        _refresh_receipt_line_stock(db, receipt_line)
    receipt_line.receipt.approval_history = [*(receipt_line.receipt.approval_history or []), {
        "action": payload.action, "discrepancy_id": str(row.id), "actor": _actor(current_user),
        "reason": payload.reason, "at": datetime.utcnow().isoformat(), "invoice_rate": str(row.invoice_rate), "po_rate": str(row.po_rate),
    }]
    row.version += 1
    db.commit()
    db.refresh(row)
    return _discrepancy_payload(row, db)


class DebitNoteCreate(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    request_id: uuid.UUID
    note_date: date
    discrepancy_ids: list[uuid.UUID] = Field(min_length=1)
    reason: str = Field(min_length=3, max_length=1000)


def _debit_payload(note: PurchaseDebitNote) -> dict[str, Any]:
    return {"id": str(note.id), "debit_note_no": note.debit_note_no, "supplier_id": str(note.supplier_id),
        "supplier_name": note.supplier_name_snapshot, "note_date": note.note_date.isoformat(), "reason": note.reason,
        "status": note.status, "version": note.version, "total_amount": float(note.total_amount),
        "settled_amount": float(note.settled_amount), "open_amount": float(note.total_amount - note.settled_amount),
        "created_by": note.created_by, "approved_by": note.approved_by,
        "lines": [{"id": str(line.id), "discrepancy_id": str(line.discrepancy_id),
                   "claimed_amount": float(line.claimed_amount), "tax_adjustment": float(line.tax_adjustment)} for line in note.lines or []]}


@router.post("/debit-notes")
def create_debit_note(
    payload: DebitNoteCreate, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
    current_user: dict = Depends(require_role(["Accounts", "PlantManager"])),
):
    request_fingerprint = canonical_hash(payload.model_dump(mode="json", exclude={"request_id"}))
    existing = db.query(PurchaseDebitNote).filter(PurchaseDebitNote.plant_id == plant_id, PurchaseDebitNote.request_id == payload.request_id).first()
    if existing:
        if existing.request_fingerprint != request_fingerprint:
            raise HTTPException(status_code=409, detail="This debit-note request key already exists with different details")
        return _debit_payload(existing)
    rows = db.query(PurchaseDiscrepancy).filter(PurchaseDiscrepancy.plant_id == plant_id, PurchaseDiscrepancy.id.in_(payload.discrepancy_ids)).with_for_update().all()
    if len(rows) != len(set(payload.discrepancy_ids)):
        raise HTTPException(status_code=404, detail="One or more discrepancies were not found")
    if any(row.claimable_amount <= 0 or row.status not in {"OPEN", "UNDER_REVIEW"} for row in rows):
        raise HTTPException(status_code=409, detail="Only open positive claim amounts can create a debit note")
    source = [_discrepancy_payload(row, db) for row in rows]
    supplier_ids = {row["supplier_id"] for row in source}
    if len(supplier_ids) != 1:
        raise HTTPException(status_code=422, detail="A debit note cannot mix vendors")
    note = PurchaseDebitNote(plant_id=plant_id, request_id=payload.request_id, request_fingerprint=request_fingerprint,
        debit_note_no=_next_doc_no(db, PurchaseDebitNote, plant_id, "debit_note_no", "DN"),
        supplier_id=uuid.UUID(next(iter(supplier_ids))), supplier_name_snapshot=source[0]["supplier_name"],
        note_date=payload.note_date, reason=payload.reason, status="DRAFT",
        total_amount=sum((row.claimable_amount for row in rows), Decimal("0")), created_by=_actor(current_user))
    db.add(note)
    db.flush()
    for row in rows:
        db.add(PurchaseDebitNoteLine(debit_note_id=note.id, discrepancy_id=row.id, claimed_amount=row.claimable_amount))
        row.status = "CLAIM_DRAFTED"
        row.version += 1
    db.commit()
    db.refresh(note)
    return _debit_payload(note)


@router.get("/debit-notes")
def list_debit_notes(db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    rows = db.query(PurchaseDebitNote).filter(PurchaseDebitNote.plant_id == plant_id).order_by(PurchaseDebitNote.created_at.desc()).all()
    return {"items": [_debit_payload(row) for row in rows]}


@router.get("/debit-notes/{note_id}/pdf")
def debit_note_pdf(
    note_id: uuid.UUID,
    db: Session = Depends(get_db),
    plant_id: str = Depends(get_current_plant),
    current_user: dict = Depends(get_current_user),
):
    note = db.query(PurchaseDebitNote).filter(
        PurchaseDebitNote.id == note_id,
        PurchaseDebitNote.plant_id == plant_id,
    ).first()
    if not note:
        raise HTTPException(status_code=404, detail="Debit note not found")
    discrepancies = []
    for line in note.lines or []:
        row = db.query(PurchaseDiscrepancy).filter(PurchaseDiscrepancy.id == line.discrepancy_id).first()
        if row:
            discrepancies.append(_discrepancy_payload(row, db))
    pdf_bytes = render_debit_note_pdf(note, discrepancies)
    filename = f"{note.debit_note_no.replace('/', '-')}.pdf"
    return StreamingResponse(
        BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


class DebitNoteAction(BaseModel):
    model_config = ConfigDict(extra="forbid")
    action: str = Field(pattern="^(SUBMIT|APPROVE|ISSUE|VOID)$")
    expected_version: int = Field(gt=0)
    reason: Optional[str] = Field(default=None, max_length=1000)


@router.post("/debit-notes/{note_id}/action")
def act_on_debit_note(
    note_id: uuid.UUID, payload: DebitNoteAction, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
    current_user: dict = Depends(require_role(["Accounts", "PlantManager"])),
):
    note = db.query(PurchaseDebitNote).filter(PurchaseDebitNote.id == note_id, PurchaseDebitNote.plant_id == plant_id).with_for_update().first()
    if not note:
        raise HTTPException(status_code=404, detail="Debit note not found")
    if note.version != payload.expected_version:
        raise HTTPException(status_code=409, detail="Debit note changed; reload before deciding")
    actor = _actor(current_user)
    if payload.action == "SUBMIT" and note.status == "DRAFT":
        note.status = "SUBMITTED"
    elif payload.action == "APPROVE" and note.status == "SUBMITTED":
        if actor.strip().lower() == note.created_by.strip().lower():
            raise HTTPException(status_code=403, detail="A different person must approve the debit note")
        note.status = "APPROVED"; note.approved_by = actor
    elif payload.action == "ISSUE" and note.status == "APPROVED":
        note.status = "ISSUED"; note.issued_at = datetime.utcnow()
        for line in note.lines or []:
            discrepancy = db.query(PurchaseDiscrepancy).filter(PurchaseDiscrepancy.id == line.discrepancy_id).first()
            if discrepancy:
                discrepancy.status = "CLAIMED"; discrepancy.version += 1
    elif payload.action == "VOID" and note.status in {"DRAFT", "SUBMITTED", "APPROVED", "ISSUED"}:
        if not payload.reason:
            raise HTTPException(status_code=422, detail="Void reason is required")
        if Decimal(note.settled_amount or 0) > 0:
            raise HTTPException(status_code=409, detail="A debit note with settlement history cannot be voided; record a controlled adjustment")
        note.status = "VOID"
        for line in note.lines or []:
            discrepancy = db.query(PurchaseDiscrepancy).filter(PurchaseDiscrepancy.id == line.discrepancy_id).first()
            if discrepancy and discrepancy.status in {"CLAIM_DRAFTED", "CLAIMED"}:
                discrepancy.status = "OPEN"
                discrepancy.resolution_reason = f"Debit note {note.debit_note_no} voided: {payload.reason.strip()}"
                discrepancy.version += 1
    else:
        raise HTTPException(status_code=409, detail=f"{payload.action} is not valid from {note.status}")
    note.version += 1
    db.commit(); db.refresh(note)
    return _debit_payload(note)


class SettlementCreate(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    amount: float = Field(gt=0)
    settlement_date: date
    reference: str = Field(min_length=1, max_length=160)


@router.post("/debit-notes/{note_id}/settlements")
def settle_debit_note(
    note_id: uuid.UUID, payload: SettlementCreate, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
    current_user: dict = Depends(require_role(["Accounts", "PlantManager"])),
):
    note = db.query(PurchaseDebitNote).filter(PurchaseDebitNote.id == note_id, PurchaseDebitNote.plant_id == plant_id).with_for_update().first()
    if not note or note.status not in {"ISSUED", "PARTIALLY_SETTLED"}:
        raise HTTPException(status_code=409, detail="Only issued debit notes can be settled")
    amount = Decimal(str(payload.amount))
    if Decimal(note.settled_amount) + amount > Decimal(note.total_amount):
        raise HTTPException(status_code=422, detail="Settlement exceeds the open debit-note amount")
    db.add(PurchaseDebitNoteSettlement(debit_note_id=note.id, amount=amount,
        settlement_date=payload.settlement_date, reference=payload.reference, created_by=_actor(current_user)))
    note.settled_amount = Decimal(note.settled_amount) + amount
    note.status = "SETTLED" if note.settled_amount == note.total_amount else "PARTIALLY_SETTLED"
    note.version += 1
    db.commit(); db.refresh(note)
    return _debit_payload(note)


class PlanEntryInput(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    id: Optional[uuid.UUID] = None
    entry_date: date
    item_id: uuid.UUID
    supplier_id: Optional[uuid.UUID] = None
    supplier_name: Optional[str] = Field(default=None, max_length=200)
    material_form: str = Field(default="REEL", pattern="^(REEL|COIL|BULK)$")
    qty_kg: float = Field(ge=0)
    expected_unit_count: Optional[int] = Field(default=None, gt=0)
    notes: Optional[str] = Field(default=None, max_length=1000)


class PlanCreate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    request_id: uuid.UUID
    month: date
    name: str = Field(min_length=2, max_length=160)
    target_mode: str = Field(default="ARRIVAL", pattern="^(ARRIVAL|CONSUMPTION|STOCK_TARGET)$")
    source_hash: Optional[str] = Field(default=None, max_length=64)
    working_calendar: dict[str, Any] = Field(default_factory=dict)
    entries: list[PlanEntryInput] = Field(default_factory=list)


class PlanUpdate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    expected_version: int = Field(gt=0)
    entries: list[PlanEntryInput]
    source_hash: Optional[str] = Field(default=None, min_length=64, max_length=64)
    source_metadata: Optional[dict[str, Any]] = None


def _plan_payload(plan: ProcurementPlan) -> dict[str, Any]:
    return {"id": str(plan.id), "month": plan.month.isoformat(), "name": plan.name, "status": plan.status,
        "target_mode": plan.target_mode, "source_hash": plan.source_hash, "working_calendar": plan.working_calendar or {},
        "version": plan.version, "created_by": plan.created_by,
        "entries": [{"id": str(entry.id), "entry_date": entry.entry_date.isoformat(), "item_id": str(entry.item_id),
            "item_code": entry.item.item_code if entry.item else None, "item_name": entry.item.name if entry.item else None,
            "supplier_id": str(entry.supplier_id) if entry.supplier_id else None, "supplier_name": entry.supplier_name_snapshot,
            "material_form": entry.material_form, "qty_kg": float(entry.qty_kg),
            "expected_unit_count": entry.expected_unit_count, "converted_qty_kg": float(entry.converted_qty_kg),
            "status": entry.status, "notes": entry.notes} for entry in plan.entries or []]}


def _replace_plan_entries(db: Session, plan: ProcurementPlan, entries: list[PlanEntryInput], plant_id: str) -> None:
    existing = {entry.id: entry for entry in plan.entries or []}
    touched: set[uuid.UUID] = set()
    for index, requested in enumerate(entries, start=1):
        if month_start(requested.entry_date) != plan.month:
            raise HTTPException(status_code=422, detail=f"Plan entry {index} is outside {plan.month.strftime('%B %Y')}")
        item = db.query(ItemMaster).filter(ItemMaster.id == requested.item_id, ItemMaster.plant_id == plant_id).first()
        if not item:
            raise HTTPException(status_code=404, detail=f"Material not found for plan entry {index}")
        if _enum(item.type) == "RAW_PAPER" and _enum(item.uom) != "KG":
            raise HTTPException(status_code=409, detail=f"Paper material {item.item_code} is not configured in KG")
        entry = existing.get(requested.id) if requested.id else None
        if entry and Decimal(entry.converted_qty_kg or 0) > Decimal(str(requested.qty_kg)):
            raise HTTPException(status_code=409, detail=f"Entry {index} is already converted above the requested quantity")
        if not entry:
            entry = ProcurementPlanEntry(plan_id=plan.id, entry_date=requested.entry_date, item_id=item.id, qty_kg=requested.qty_kg)
            db.add(entry); db.flush()
        entry.entry_date = requested.entry_date; entry.item_id = item.id; entry.supplier_id = requested.supplier_id
        entry.supplier_name_snapshot = requested.supplier_name; entry.material_form = requested.material_form
        entry.qty_kg = requested.qty_kg; entry.expected_unit_count = requested.expected_unit_count; entry.notes = requested.notes
        entry.status = "PLANNED" if requested.qty_kg > float(entry.converted_qty_kg or 0) else "CONVERTED"
        touched.add(entry.id)
    for entry_id, entry in existing.items():
        if entry_id not in touched:
            if Decimal(entry.converted_qty_kg or 0) > 0:
                raise HTTPException(status_code=409, detail="Converted plan entries cannot be deleted")
            db.delete(entry)


@router.get("/plans")
def list_plans(month: Optional[date] = Query(default=None), db: Session = Depends(get_db),
               plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    query = db.query(ProcurementPlan).filter(ProcurementPlan.plant_id == plant_id)
    if month:
        query = query.filter(ProcurementPlan.month == month_start(month))
    return {"items": [_plan_payload(row) for row in query.order_by(ProcurementPlan.month.desc(), ProcurementPlan.created_at.desc()).all()]}


@router.post("/plans")
def create_plan(payload: PlanCreate, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
                current_user: dict = Depends(require_role(["Planner", "Store", "PlantManager"]))):
    request_fingerprint = canonical_hash(payload.model_dump(mode="json", exclude={"request_id"}))
    existing = db.query(ProcurementPlan).filter(ProcurementPlan.plant_id == plant_id, ProcurementPlan.request_id == payload.request_id).first()
    if existing:
        if existing.request_fingerprint != request_fingerprint:
            raise HTTPException(status_code=409, detail="This plan request key already exists with different details")
        return _plan_payload(existing)
    plan = ProcurementPlan(plant_id=plant_id, request_id=payload.request_id, request_fingerprint=request_fingerprint, month=month_start(payload.month),
        name=payload.name.strip(), target_mode=payload.target_mode, source_hash=payload.source_hash,
        working_calendar=payload.working_calendar, created_by=_actor(current_user))
    db.add(plan)
    try:
        db.flush()
    except IntegrityError as error:
        db.rollback()
        # Two tabs can create the same month before either GET finishes.
        # Preserve idempotent replays and report a recoverable conflict.
        prior = db.query(ProcurementPlan).filter(
            ProcurementPlan.plant_id == plant_id,
            ProcurementPlan.request_id == payload.request_id,
        ).first()
        if prior and prior.request_fingerprint == request_fingerprint:
            return _plan_payload(prior)
        constraint = getattr(getattr(error.orig, "diag", None), "constraint_name", "")
        if prior or constraint == "uq_procurement_plan_month_name":
            raise HTTPException(status_code=409, detail="A plan with this name already exists for the month. Refresh and select the saved plan.") from error
        raise
    _replace_plan_entries(db, plan, payload.entries, plant_id)
    db.commit(); db.refresh(plan)
    return _plan_payload(plan)


@router.put("/plans/{plan_id}/entries")
def update_plan_entries(plan_id: uuid.UUID, payload: PlanUpdate, db: Session = Depends(get_db),
                        plant_id: str = Depends(get_current_plant), current_user: dict = Depends(require_role(["Planner", "Store", "PlantManager"]))):
    plan = db.query(ProcurementPlan).filter(ProcurementPlan.id == plan_id, ProcurementPlan.plant_id == plant_id).with_for_update().first()
    if not plan:
        raise HTTPException(status_code=404, detail="Procurement plan not found")
    if plan.version != payload.expected_version:
        raise HTTPException(status_code=409, detail="Calendar changed; reload before saving")
    if plan.status not in {"DRAFT", "REJECTED"}:
        raise HTTPException(status_code=409, detail="Only a draft plan can be edited")
    _replace_plan_entries(db, plan, payload.entries, plant_id)
    if payload.source_hash:
        plan.source_hash = payload.source_hash
    if payload.source_metadata is not None:
        plan.working_calendar = {**(plan.working_calendar or {}), "source_import": payload.source_metadata}
    plan.version += 1
    db.commit(); db.refresh(plan)
    return _plan_payload(plan)


class PlanStatusAction(BaseModel):
    model_config = ConfigDict(extra="forbid")
    action: str = Field(pattern="^(SUBMIT|APPROVE|REJECT|LOCK)$")
    expected_version: int = Field(gt=0)
    reason: Optional[str] = Field(default=None, max_length=1000)


@router.post("/plans/{plan_id}/action")
def act_on_plan(plan_id: uuid.UUID, payload: PlanStatusAction, db: Session = Depends(get_db),
                plant_id: str = Depends(get_current_plant), current_user: dict = Depends(require_role(["Planner", "PlantManager"]))):
    plan = db.query(ProcurementPlan).filter(ProcurementPlan.id == plan_id, ProcurementPlan.plant_id == plant_id).with_for_update().first()
    if not plan:
        raise HTTPException(status_code=404, detail="Procurement plan not found")
    if plan.version != payload.expected_version:
        raise HTTPException(status_code=409, detail="Plan changed; reload before deciding")
    actor = _actor(current_user)
    if payload.action == "SUBMIT" and plan.status in {"DRAFT", "REJECTED"}:
        plan.status = "SUBMITTED"
    elif payload.action == "APPROVE" and plan.status == "SUBMITTED":
        if actor.strip().lower() == plan.created_by.strip().lower():
            raise HTTPException(status_code=403, detail="A different person must approve the procurement plan")
        plan.status = "APPROVED"
    elif payload.action == "REJECT" and plan.status == "SUBMITTED":
        if not payload.reason:
            raise HTTPException(status_code=422, detail="Rejection reason is required")
        plan.status = "REJECTED"
    elif payload.action == "LOCK" and plan.status == "APPROVED":
        plan.status = "LOCKED"
    else:
        raise HTTPException(status_code=409, detail=f"{payload.action} is not valid from {plan.status}")
    plan.version += 1
    db.commit(); db.refresh(plan)
    return _plan_payload(plan)


class PlanConversionInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    request_id: uuid.UUID
    expected_plan_version: int = Field(gt=0)
    entry_ids: list[uuid.UUID] = Field(min_length=1)


def _effective_rm_cost(db: Session, sheet: Optional[RmCostSheet], as_of: date) -> Optional[RmCostVersion]:
    if not sheet:
        return None
    return db.query(RmCostVersion).filter(
        RmCostVersion.sheet_id == sheet.id,
        RmCostVersion.status.in_(["ACTIVE", "SCHEDULED"]),
        RmCostVersion.effective_from <= as_of,
        (RmCostVersion.effective_to.is_(None)) | (RmCostVersion.effective_to > as_of),
    ).order_by(RmCostVersion.effective_from.desc(), RmCostVersion.version_no.desc()).first()


@router.post("/plans/{plan_id}/convert")
def convert_plan_to_purchase_orders(plan_id: uuid.UUID, payload: PlanConversionInput, db: Session = Depends(get_db),
                                    plant_id: str = Depends(get_current_plant), current_user: dict = Depends(require_role(["Planner", "Store"]))):
    request_hash = canonical_hash(payload.model_dump(mode="json", exclude={"request_id"}))
    existing = db.query(PlanConversion).filter(PlanConversion.plant_id == plant_id, PlanConversion.request_id == payload.request_id).first()
    if existing:
        if existing.payload_hash != request_hash:
            raise HTTPException(status_code=409, detail="Conversion request key already exists with different entries")
        return {"id": str(existing.id), "purchase_order_ids": existing.purchase_order_ids, "idempotent": True}
    plan = db.query(ProcurementPlan).filter(ProcurementPlan.id == plan_id, ProcurementPlan.plant_id == plant_id).with_for_update().first()
    if not plan or plan.status not in {"APPROVED", "LOCKED"}:
        raise HTTPException(status_code=409, detail="Only an approved procurement plan can generate POs")
    if plan.version != payload.expected_plan_version:
        raise HTTPException(status_code=409, detail="Plan changed; run conversion preview again")
    entries = db.query(ProcurementPlanEntry).filter(ProcurementPlanEntry.plan_id == plan.id, ProcurementPlanEntry.id.in_(payload.entry_ids)).with_for_update().all()
    if len(entries) != len(set(payload.entry_ids)):
        raise HTTPException(status_code=404, detail="One or more plan entries were not found")
    groups: dict[tuple[str, str], list[ProcurementPlanEntry]] = {}
    for entry in entries:
        uncovered = Decimal(entry.qty_kg) - Decimal(entry.converted_qty_kg or 0)
        if uncovered <= 0:
            continue
        if not entry.supplier_id or not entry.supplier_name_snapshot:
            raise HTTPException(status_code=422, detail=f"Assign a vendor before converting {entry.item.item_code}")
        groups.setdefault((str(entry.supplier_id), entry.supplier_name_snapshot), []).append(entry)
    if not groups:
        raise HTTPException(status_code=409, detail="Selected entries are already fully converted")
    order_ids: list[str] = []
    for (supplier_id, supplier_name), grouped in groups.items():
        order = PurchaseOrder(plant_id=plant_id, request_id=uuid.uuid4(), po_no=_next_purchase_order_no(db, plant_id, "RM_PM"),
            request_fingerprint=request_hash, category="RM_PM", current_revision_no=1, version=1,
            supplier_id=uuid.UUID(supplier_id), supplier_name_snapshot=supplier_name,
            expected_date=min(entry.entry_date for entry in grouped), status="DRAFT",
            notes=f"Generated from procurement plan {plan.name}",
            metadata_json={"po_date": date.today().isoformat(), "source_plan_id": str(plan.id)}, created_by=_actor(current_user))
        db.add(order); db.flush()
        po_lines: list[PurchaseOrderLine] = []
        line_entries: list[tuple[PurchaseOrderLine, ProcurementPlanEntry, Decimal]] = []
        for entry in grouped:
            uncovered = Decimal(entry.qty_kg) - Decimal(entry.converted_qty_kg or 0)
            cost_sheet = db.query(RmCostSheet).filter(
                RmCostSheet.plant_id == plant_id,
                RmCostSheet.item_id == entry.item_id,
                RmCostSheet.currency == "INR",
            ).first()
            active_cost = _effective_rm_cost(db, cost_sheet, date.today())
            if not active_cost:
                raise HTTPException(status_code=422, detail=f"{entry.item.item_code} has no active INR planning cost; set RM costing before PO conversion")
            line = PurchaseOrderLine(purchase_order_id=order.id, logical_line_id=uuid.uuid4(), item_id=entry.item_id,
                qty_ordered=float(uncovered), qty_received=0, unit_cost=float(active_cost.landed_cost), uom="KG",
                expected_unit_count=entry.expected_unit_count, count_basis="ESTIMATED" if entry.expected_unit_count else None,
                line_status="OPEN", incoming_qc_required=True,
                metadata_json={"description": entry.item.name, "source_plan_entry_id": str(entry.id),
                    "delivery_date": entry.entry_date.isoformat(), "draft_rate_source": "RM_STANDARD_COST",
                    "draft_cost_version_id": str(active_cost.id), "draft_cost_version": active_cost.version_no})
            db.add(line); po_lines.append(line)
            line_entries.append((line, entry, uncovered))
            entry.converted_qty_kg = Decimal(entry.converted_qty_kg or 0) + uncovered
            entry.status = "CONVERTED"
        db.flush()
        revision = _persist_revision(db, order, po_lines, revision_no=1, request_id=order.request_id, actor=_actor(current_user))
        db.flush()
        revision_line_by_source = {row.source_order_line_id: row for row in revision.lines or []}
        for line, entry, uncovered in line_entries:
            revision_line = revision_line_by_source.get(line.id)
            if not revision_line:
                raise HTTPException(status_code=500, detail="Could not create delivery allocation for generated PO line")
            db.add(PurchaseDeliverySchedule(
                revision_line_id=revision_line.id,
                purchase_order_line_id=line.id,
                plant_id=plant_id,
                created_by=_actor(current_user),
                promised_date=entry.entry_date,
                delivery_date=entry.entry_date,
                planned_qty=uncovered,
                cancelled_qty=0,
                confirmation_status="TENTATIVE",
                source_plan_entry_id=entry.id,
            ))
        order_ids.append(str(order.id))
    conversion = PlanConversion(plant_id=plant_id, request_id=payload.request_id, plan_id=plan.id,
        entry_ids=[str(value) for value in payload.entry_ids], purchase_order_ids=order_ids,
        payload_hash=request_hash, created_by=_actor(current_user))
    db.add(conversion); plan.version += 1
    db.commit()
    return {"id": str(conversion.id), "purchase_order_ids": order_ids, "idempotent": False}


class MrpDatedDemand(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    date: date
    qty_kg: float = Field(ge=0)


class MrpItemInput(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    item_id: uuid.UUID
    demand_kg: float = Field(ge=0)
    dated_demand: list[MrpDatedDemand] = Field(default_factory=list)
    target_stock_kg: Optional[float] = Field(default=None, ge=0)


class MrpRunInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    as_of_date: date
    horizon_end: date
    demand_source_version: str = Field(min_length=1, max_length=160)
    items: list[MrpItemInput] = Field(min_length=1)


@router.post("/mrp-runs")
def run_mrp(payload: MrpRunInput, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
            current_user: dict = Depends(require_role(["Planner", "Store", "PlantManager"]))):
    if payload.horizon_end < payload.as_of_date:
        raise HTTPException(status_code=422, detail="MRP horizon cannot end before the as-of date")
    results = []
    policy_versions: dict[str, int] = {}
    cost_versions: dict[str, int] = {}
    for requested in payload.items:
        item = db.query(ItemMaster).filter(ItemMaster.id == requested.item_id, ItemMaster.plant_id == plant_id).first()
        if not item:
            raise HTTPException(status_code=404, detail=f"MRP material not found: {requested.item_id}")
        opening = max(0, get_usable_item_qty(str(item.id), db))
        supply = []
        po_lines = db.query(PurchaseOrderLine).join(PurchaseOrder).filter(
            PurchaseOrder.plant_id == plant_id, PurchaseOrderLine.item_id == item.id,
            PurchaseOrder.status.in_(["APPROVED", "PARTIALLY_RECEIVED"])).all()
        undated_supply = 0.0
        unconfirmed_supply = 0.0
        overdue_supply = Decimal(0)
        for po_line in po_lines:
            remaining = max(Decimal(0), Decimal(str(po_line.qty_ordered)) - Decimal(str(po_line.qty_received)) - Decimal(str(po_line.qty_short_closed or 0)))
            schedules = db.query(PurchaseDeliverySchedule).filter(
                PurchaseDeliverySchedule.purchase_order_line_id == po_line.id,
                PurchaseDeliverySchedule.confirmation_status != "CANCELLED"
            ).order_by(PurchaseDeliverySchedule.delivery_date).all()
            for schedule in schedules:
                qty = min(remaining, max(Decimal(0), Decimal(str(schedule.planned_qty)) - Decimal(str(schedule.received_qty or 0)) - Decimal(str(schedule.cancelled_qty or 0))))
                if qty <= 0:
                    continue
                remaining -= qty
                if schedule.confirmation_status != "CONFIRMED":
                    unconfirmed_supply += float(qty)
                elif schedule.delivery_date < payload.as_of_date:
                    # A missed promise is not stock arriving today. It must be re-confirmed.
                    overdue_supply += qty
                else:
                    supply.append({"date": schedule.delivery_date, "qty": qty, "po_no": po_line.order.po_no})
            undated_supply += float(remaining)
        committed = sum((row["qty"] for row in supply if row["date"] <= payload.horizon_end), Decimal(0))
        if requested.dated_demand:
            if any(row.date < payload.as_of_date or row.date > payload.horizon_end for row in requested.dated_demand):
                raise HTTPException(422, "Dated demand must fall inside the MRP horizon")
            if abs(sum(Decimal(str(row.qty_kg)) for row in requested.dated_demand) - Decimal(str(requested.demand_kg))) > Decimal("0.001"):
                raise HTTPException(422, "Daily material requirements must reconcile to total demand")
        projection = Decimal(str(opening)); timeline = []
        daily = {}
        for row in requested.dated_demand:
            daily[row.date] = daily.get(row.date, Decimal(0)) + Decimal(str(row.qty_kg))
        supply_dates = {}
        for row in supply:
            day = max(payload.as_of_date, row["date"])
            if day <= payload.horizon_end:
                supply_dates[day] = supply_dates.get(day, Decimal(0)) + row["qty"]
        for day in sorted(set(daily) | set(supply_dates)):
            projection += supply_dates.get(day, Decimal(0)) - daily.get(day, Decimal(0))
            timeline.append({"date": day.isoformat(), "demand_kg": float(daily.get(day, 0)), "committed_arrival_kg": float(supply_dates.get(day, 0)), "projected_free_kg": float(projection), "shortfall_kg": float(max(Decimal(0), -projection))})
        policy = db.query(StockAlertPolicy).filter(StockAlertPolicy.plant_id == plant_id,
            StockAlertPolicy.item_id == item.id, StockAlertPolicy.status == "ACTIVE").order_by(StockAlertPolicy.version.desc()).first()
        target = requested.target_stock_kg if requested.target_stock_kg is not None else float(policy.target_stock_kg if policy else item.reorder_level or 0)
        calculation = mrp_suggestion(opening_stock_kg=opening, demand_kg=requested.demand_kg,
            committed_supply_kg=committed, target_stock_kg=target,
            minimum_order_kg=policy.minimum_order_kg if policy else 0,
            order_multiple_kg=policy.order_multiple_kg if policy else 0)
        cost_sheet = db.query(RmCostSheet).filter(RmCostSheet.plant_id == plant_id, RmCostSheet.item_id == item.id).first()
        cost_version = _effective_rm_cost(db, cost_sheet, payload.as_of_date)
        if policy: policy_versions[str(item.id)] = policy.version
        if cost_version: cost_versions[str(item.id)] = cost_version.version_no
        results.append({"item_id": str(item.id), "item_code": item.item_code, "item_name": item.name,
            **{key: float(value) for key, value in calculation.items()},
            "daily_projection": timeline,
            "first_shortage_date": next((row["date"] for row in timeline if row["shortfall_kg"] > 0), None),
            "peak_timing_shortfall_kg": max((row["shortfall_kg"] for row in timeline), default=0),
            "undated_open_supply_kg": undated_supply,
            "unconfirmed_supply_kg": unconfirmed_supply,
            "overdue_supply_kg": float(overdue_supply),
            "estimated_rate": float(cost_version.landed_cost) if cost_version else None,
            "estimated_value": float(Decimal(str(calculation["suggested_order_kg"])) * Decimal(cost_version.landed_cost)) if cost_version else None,
            "cost_status": "AVAILABLE" if cost_version else "MISSING",
            "cost_version": cost_version.version_no if cost_version else None,
            "explanation": "Demand + target - free stock - approved open PO, then MOQ/order multiple."})
    run = MrpRun(plant_id=plant_id, as_of_date=payload.as_of_date, horizon_end=payload.horizon_end,
        source_versions={"demand": payload.demand_source_version, "policies": policy_versions, "costs": cost_versions},
        results_json=results, created_by=_actor(current_user))
    db.add(run); db.commit(); db.refresh(run)
    return {"id": str(run.id), "as_of_date": run.as_of_date.isoformat(), "horizon_end": run.horizon_end.isoformat(),
        "status": run.status, "source_versions": run.source_versions, "results": run.results_json}


@router.get("/mrp-runs/{run_id}")
def get_mrp_run(run_id: uuid.UUID, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    run = db.query(MrpRun).filter(MrpRun.id == run_id, MrpRun.plant_id == plant_id).first()
    if not run:
        raise HTTPException(status_code=404, detail="MRP run not found")
    return {"id": str(run.id), "as_of_date": run.as_of_date.isoformat(), "horizon_end": run.horizon_end.isoformat(),
        "status": run.status, "source_versions": run.source_versions, "results": run.results_json,
        "created_at": run.created_at.isoformat(), "created_by": run.created_by}


class StockPolicyCreate(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    request_id: uuid.UUID
    item_id: uuid.UUID
    stock_basis: str = Field(default="FREE_STOCK", pattern="^(FREE_STOCK|PHYSICAL_STOCK)$")
    safety_stock_kg: float = Field(ge=0)
    reorder_point_kg: float = Field(ge=0)
    target_stock_kg: float = Field(ge=0)
    recovery_margin_kg: float = Field(default=0, ge=0)
    lead_time_days: int = Field(default=0, ge=0)
    minimum_order_kg: float = Field(default=0, ge=0)
    order_multiple_kg: float = Field(default=0, ge=0)
    recipients: list[str] = Field(default_factory=list)
    cooldown_hours: int = Field(default=24, ge=1, le=720)
    change_reason: str = Field(min_length=3, max_length=1000)

    @field_validator("target_stock_kg")
    @classmethod
    def finite_target(cls, value: float) -> float:
        return value


def _policy_payload(policy: StockAlertPolicy) -> dict[str, Any]:
    return {"id": str(policy.id), "item_id": str(policy.item_id) if policy.item_id else None,
        "item_code": policy.item.item_code if getattr(policy, "item", None) else None,
        "item_name": policy.item.name if getattr(policy, "item", None) else None,
        "scope_type": policy.scope_type, "status": policy.status, "stock_basis": policy.stock_basis,
        "safety_stock_kg": float(policy.safety_stock_kg), "reorder_point_kg": float(policy.reorder_point_kg),
        "target_stock_kg": float(policy.target_stock_kg), "recovery_margin_kg": float(policy.recovery_margin_kg),
        "lead_time_days": policy.lead_time_days, "minimum_order_kg": float(policy.minimum_order_kg),
        "order_multiple_kg": float(policy.order_multiple_kg), "recipients": policy.recipients or [],
        "cooldown_hours": policy.cooldown_hours, "change_reason": policy.change_reason,
        "activation_reason": policy.activation_reason, "version": policy.version, "created_by": policy.created_by,
        "activated_by": policy.activated_by, "activated_at": policy.activated_at.isoformat() if policy.activated_at else None}


@router.get("/stock-alert-policies")
def list_stock_policies(db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    rows = db.query(StockAlertPolicy).filter(StockAlertPolicy.plant_id == plant_id).order_by(StockAlertPolicy.item_id, StockAlertPolicy.version.desc()).all()
    return {"items": [_policy_payload(row) for row in rows]}


@router.post("/stock-alert-policies")
def create_stock_policy(payload: StockPolicyCreate, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
                        current_user: dict = Depends(require_role(["PlantManager", "Store"]))):
    if not (payload.target_stock_kg >= payload.reorder_point_kg >= payload.safety_stock_kg):
        raise HTTPException(status_code=422, detail="Target stock must be at least reorder point, which must be at least safety stock")
    request_fingerprint = canonical_hash(payload.model_dump(mode="json", exclude={"request_id"}))
    existing_request = db.query(StockAlertPolicy).filter(StockAlertPolicy.plant_id == plant_id, StockAlertPolicy.request_id == payload.request_id).first()
    if existing_request:
        if existing_request.request_fingerprint != request_fingerprint:
            raise HTTPException(status_code=409, detail="This policy request key already exists with different details")
        return _policy_payload(existing_request)
    item = db.query(ItemMaster).filter(ItemMaster.id == payload.item_id, ItemMaster.plant_id == plant_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Material not found")
    if _enum(item.type) == "RAW_PAPER" and _enum(item.uom) != "KG":
        raise HTTPException(status_code=409, detail="Paper stock policy requires a KG material master")
    latest = db.query(func.coalesce(func.max(StockAlertPolicy.version), 0)).filter(
        StockAlertPolicy.plant_id == plant_id, StockAlertPolicy.item_id == item.id).scalar() or 0
    policy = StockAlertPolicy(plant_id=plant_id, request_id=payload.request_id, request_fingerprint=request_fingerprint, item_id=item.id, scope_type="ITEM",
        status="DRAFT", stock_basis=payload.stock_basis, safety_stock_kg=payload.safety_stock_kg,
        reorder_point_kg=payload.reorder_point_kg, target_stock_kg=payload.target_stock_kg,
        recovery_margin_kg=payload.recovery_margin_kg, lead_time_days=payload.lead_time_days,
        minimum_order_kg=payload.minimum_order_kg, order_multiple_kg=payload.order_multiple_kg,
        recipients=payload.recipients, cooldown_hours=payload.cooldown_hours, version=int(latest) + 1,
        change_reason=payload.change_reason.strip(), created_by=_actor(current_user))
    db.add(policy); db.commit(); db.refresh(policy)
    return _policy_payload(policy)


class VersionAction(BaseModel):
    model_config = ConfigDict(extra="forbid")
    expected_version: int = Field(gt=0)
    reason: str = Field(min_length=3, max_length=1000)


@router.post("/stock-alert-policies/{policy_id}/activate")
def activate_stock_policy(policy_id: uuid.UUID, payload: VersionAction, db: Session = Depends(get_db),
                          plant_id: str = Depends(get_current_plant), current_user: dict = Depends(require_role(["PlantManager"]))):
    policy = db.query(StockAlertPolicy).filter(StockAlertPolicy.id == policy_id, StockAlertPolicy.plant_id == plant_id).with_for_update().first()
    if not policy:
        raise HTTPException(status_code=404, detail="Stock alert policy not found")
    if policy.version != payload.expected_version or policy.status != "DRAFT":
        raise HTTPException(status_code=409, detail="Policy changed or is no longer a draft")
    if policy.created_by.strip().lower() == _actor(current_user).strip().lower():
        raise HTTPException(status_code=403, detail="A different person must activate the stock alert policy")
    db.query(StockAlertPolicy).filter(StockAlertPolicy.plant_id == plant_id,
        StockAlertPolicy.item_id == policy.item_id, StockAlertPolicy.status == "ACTIVE").update({"status": "INACTIVE"})
    policy.status = "ACTIVE"; policy.activation_reason = payload.reason.strip()
    policy.activated_by = _actor(current_user); policy.activated_at = datetime.utcnow()
    # Keep legacy fields aligned so old MRP readers do not disagree.
    item = db.query(ItemMaster).filter(ItemMaster.id == policy.item_id).first()
    if item:
        item.safety_stock = float(policy.safety_stock_kg); item.reorder_level = float(policy.reorder_point_kg)
        item.lead_time_days = float(policy.lead_time_days)
    db.commit(); db.refresh(policy)
    return _policy_payload(policy)


def _alert_payload(row: StockAlertEpisode, db: Session) -> dict[str, Any]:
    item = db.query(ItemMaster).filter(ItemMaster.id == row.item_id).first()
    policy = db.query(StockAlertPolicy).filter(StockAlertPolicy.id == row.policy_id).first()
    return {"id": str(row.id), "policy_id": str(row.policy_id), "policy_version": policy.version if policy else None,
        "item_id": str(row.item_id), "item_code": item.item_code if item else None, "item_name": item.name if item else None,
        "status": row.status, "severity": row.severity, "stock_qty_kg": float(row.stock_qty_kg),
        "threshold_qty_kg": float(row.threshold_qty_kg), "assignee": row.assignee,
        "acknowledged_by": row.acknowledged_by, "acknowledged_at": row.acknowledged_at.isoformat() if row.acknowledged_at else None,
        "snoozed_until": row.snoozed_until.isoformat() if row.snoozed_until else None,
        "breached_at": row.breached_at.isoformat(), "recovered_at": row.recovered_at.isoformat() if row.recovered_at else None}


@router.post("/stock-alerts/evaluate")
def evaluate_stock_alerts(db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
                          current_user: dict = Depends(require_role(["PlantManager", "Store", "Planner"]))):
    policies = db.query(StockAlertPolicy).filter(StockAlertPolicy.plant_id == plant_id, StockAlertPolicy.status == "ACTIVE").all()
    changed: list[StockAlertEpisode] = []
    for policy in policies:
        stock = max(0, get_usable_item_qty(str(policy.item_id), db)) if policy.stock_basis == "FREE_STOCK" else max(0, float(get_usable_item_qty(str(policy.item_id), db)))
        severity = alert_severity(stock, safety_stock_kg=policy.safety_stock_kg, reorder_point_kg=policy.reorder_point_kg)
        open_episode = db.query(StockAlertEpisode).filter(StockAlertEpisode.plant_id == plant_id,
            StockAlertEpisode.item_id == policy.item_id, StockAlertEpisode.status.in_(["OPEN", "ACKNOWLEDGED", "SNOOZED"])).first()
        if severity and not open_episode:
            recent_recovery = db.query(StockAlertEpisode).filter(
                StockAlertEpisode.plant_id == plant_id,
                StockAlertEpisode.item_id == policy.item_id,
                StockAlertEpisode.status == "RECOVERED",
                StockAlertEpisode.recovered_at >= datetime.utcnow() - timedelta(hours=policy.cooldown_hours),
            ).first()
            if recent_recovery:
                continue
            open_episode = StockAlertEpisode(plant_id=plant_id, policy_id=policy.id, item_id=policy.item_id,
                severity=severity, stock_qty_kg=stock,
                threshold_qty_kg=policy.safety_stock_kg if severity == "CRITICAL" else policy.reorder_point_kg)
            db.add(open_episode); changed.append(open_episode)
        elif severity and open_episode:
            open_episode.severity = severity; open_episode.stock_qty_kg = stock
            open_episode.threshold_qty_kg = policy.safety_stock_kg if severity == "CRITICAL" else policy.reorder_point_kg
        elif open_episode and Decimal(str(stock)) > Decimal(policy.reorder_point_kg) + Decimal(policy.recovery_margin_kg):
            open_episode.status = "RECOVERED"; open_episode.stock_qty_kg = stock; open_episode.recovered_at = datetime.utcnow()
            changed.append(open_episode)
    db.commit()
    return {"evaluated_policies": len(policies), "changed": [_alert_payload(row, db) for row in changed]}


@router.get("/stock-alerts")
def list_stock_alerts(status: Optional[str] = Query(default=None), db: Session = Depends(get_db),
                      plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    query = db.query(StockAlertEpisode).filter(StockAlertEpisode.plant_id == plant_id)
    if status: query = query.filter(StockAlertEpisode.status == status.strip().upper())
    rows = query.order_by(StockAlertEpisode.breached_at.desc()).all()
    return {"items": [_alert_payload(row, db) for row in rows]}


class AlertAction(BaseModel):
    model_config = ConfigDict(extra="forbid")
    action: str = Field(pattern="^(ACKNOWLEDGE|SNOOZE|ASSIGN)$")
    assignee: Optional[str] = Field(default=None, max_length=200)
    snoozed_until: Optional[datetime] = None


@router.post("/stock-alerts/{alert_id}/action")
def act_on_alert(alert_id: uuid.UUID, payload: AlertAction, db: Session = Depends(get_db),
                 plant_id: str = Depends(get_current_plant), current_user: dict = Depends(require_role(["Planner", "Store", "PlantManager"]))):
    row = db.query(StockAlertEpisode).filter(StockAlertEpisode.id == alert_id, StockAlertEpisode.plant_id == plant_id).first()
    if not row or row.status == "RECOVERED":
        raise HTTPException(status_code=409, detail="Open stock alert not found")
    if payload.action == "ACKNOWLEDGE":
        row.status = "ACKNOWLEDGED"; row.acknowledged_by = _actor(current_user); row.acknowledged_at = datetime.utcnow()
    elif payload.action == "SNOOZE":
        if not payload.snoozed_until or payload.snoozed_until <= datetime.utcnow():
            raise HTTPException(status_code=422, detail="Snooze must end in the future")
        row.status = "SNOOZED"; row.snoozed_until = payload.snoozed_until
    elif payload.action == "ASSIGN":
        if not payload.assignee: raise HTTPException(status_code=422, detail="Assignee is required")
        row.assignee = payload.assignee.strip()
    db.commit(); db.refresh(row)
    return _alert_payload(row, db)


class CostComponentInput(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    component_type: str = Field(min_length=2, max_length=30)
    label: str = Field(min_length=2, max_length=120)
    calculation_mode: str = Field(pattern="^(PER_KG|PERCENT_BASE|FIXED_PER_LOT)$")
    entered_value: float = Field(ge=0)
    reference_qty_kg: Optional[float] = Field(default=None, gt=0)


class CostDraftCreate(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    request_id: uuid.UUID
    item_id: uuid.UUID
    currency: str = Field(default="INR", pattern="^[A-Z]{3}$")
    base_cost: float = Field(ge=0)
    effective_from: date
    change_reason: str = Field(min_length=3, max_length=1000)
    source: Optional[str] = Field(default=None, max_length=80)
    expected_sheet_version: Optional[int] = Field(default=None, gt=0)
    components: list[CostComponentInput] = Field(default_factory=list)


def _cost_version_payload(version: RmCostVersion) -> dict[str, Any]:
    return {"id": str(version.id), "sheet_id": str(version.sheet_id), "version_no": version.version_no,
        "status": version.status, "base_cost": float(version.base_cost), "landed_cost": float(version.landed_cost),
        "effective_from": version.effective_from.isoformat(), "effective_to": version.effective_to.isoformat() if version.effective_to else None,
        "change_reason": version.change_reason, "activation_reason": version.activation_reason,
        "source": version.source, "content_hash": version.content_hash,
        "created_by": version.created_by, "activated_by": version.activated_by,
        "activated_at": version.activated_at.isoformat() if version.activated_at else None,
        "created_at": version.created_at.isoformat(), "components": [{"id": str(component.id),
            "component_type": component.component_type, "label": component.label,
            "calculation_mode": component.calculation_mode, "entered_value": float(component.entered_value),
            "normalized_per_kg": float(component.normalized_per_kg), "metadata": component.metadata_json or {}}
            for component in version.components or []]}


@router.get("/rm-costing")
def list_rm_costing(db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    items = db.query(ItemMaster).filter(ItemMaster.plant_id == plant_id, ItemMaster.type != "FINISHED_GOOD", ItemMaster.active == "true").order_by(ItemMaster.item_code).all()
    rows = []
    for item in items:
        sheet = db.query(RmCostSheet).filter(RmCostSheet.plant_id == plant_id, RmCostSheet.item_id == item.id).first()
        current = _effective_rm_cost(db, sheet, date.today())
        latest = db.query(RmCostVersion).filter(RmCostVersion.sheet_id == sheet.id).order_by(RmCostVersion.version_no.desc()).first() if sheet else None
        rows.append({"item_id": str(item.id), "item_code": item.item_code, "item_name": item.name, "uom": _enum(item.uom),
            "sheet_id": str(sheet.id) if sheet else None, "sheet_version": sheet.version if sheet else None,
            "currency": sheet.currency if sheet else "INR", "base_cost": float(current.base_cost) if current else None,
            "landed_cost": float(current.landed_cost) if current else None, "active_version": current.version_no if current else None,
            "latest_version": latest.version_no if latest else None, "latest_status": latest.status if latest else "MISSING",
            "effective_from": current.effective_from.isoformat() if current else None})
    return {"items": rows}


@router.get("/rm-costing/{item_id}/history")
def cost_history(item_id: uuid.UUID, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    sheet = db.query(RmCostSheet).filter(RmCostSheet.plant_id == plant_id, RmCostSheet.item_id == item_id).first()
    if not sheet: return {"item_id": str(item_id), "items": []}
    rows = db.query(RmCostVersion).filter(RmCostVersion.sheet_id == sheet.id).order_by(RmCostVersion.version_no.desc()).all()
    return {"item_id": str(item_id), "sheet_id": str(sheet.id), "sheet_version": sheet.version,
            "items": [_cost_version_payload(row) for row in rows]}


@router.post("/rm-costing/drafts")
def create_cost_draft(payload: CostDraftCreate, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
                      current_user: dict = Depends(require_cost_administrator)):
    item = db.query(ItemMaster).filter(ItemMaster.id == payload.item_id, ItemMaster.plant_id == plant_id).first()
    if not item or _enum(item.type) == "FINISHED_GOOD":
        raise HTTPException(status_code=404, detail="Raw material not found")
    request_fingerprint = canonical_hash(payload.model_dump(mode="json", exclude={"request_id", "expected_sheet_version"}))
    sheet = db.query(RmCostSheet).filter(RmCostSheet.plant_id == plant_id, RmCostSheet.item_id == item.id,
        RmCostSheet.currency == payload.currency).with_for_update().first()
    if not sheet:
        if payload.expected_sheet_version is not None:
            raise HTTPException(status_code=409, detail="Cost sheet did not exist; reload before saving")
        sheet = RmCostSheet(plant_id=plant_id, item_id=item.id, currency=payload.currency, base_uom="KG" if _enum(item.type) == "RAW_PAPER" else _enum(item.uom))
        db.add(sheet); db.flush()
    else:
        replay = db.query(RmCostVersion).filter(RmCostVersion.sheet_id == sheet.id, RmCostVersion.request_id == payload.request_id).first()
        if replay:
            if replay.request_fingerprint != request_fingerprint:
                raise HTTPException(status_code=409, detail="This cost request key already exists with different details")
            return _cost_version_payload(replay)
        if payload.expected_sheet_version != sheet.version:
            raise HTTPException(status_code=409, detail="Cost sheet changed; reload before saving")
    try:
        landed, components = calculate_landed_cost(payload.base_cost, [row.model_dump() for row in payload.components])
    except ProcurementRuleError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    latest_no = db.query(func.coalesce(func.max(RmCostVersion.version_no), 0)).filter(RmCostVersion.sheet_id == sheet.id).scalar() or 0
    snapshot = {"item_id": str(item.id), "base_cost": payload.base_cost, "landed_cost": str(landed),
        "effective_from": payload.effective_from.isoformat(), "components": components, "reason": payload.change_reason}
    version = RmCostVersion(sheet_id=sheet.id, version_no=int(latest_no) + 1, request_id=payload.request_id,
        request_fingerprint=request_fingerprint,
        status="DRAFT", base_cost=payload.base_cost, landed_cost=landed, effective_from=payload.effective_from,
        change_reason=payload.change_reason, source=payload.source, content_hash=canonical_hash(snapshot), created_by=_actor(current_user))
    db.add(version); db.flush()
    for index, component in enumerate(components):
        original = payload.components[index]
        db.add(RmCostComponent(version_id=version.id, component_type=component["component_type"], label=component["label"],
            calculation_mode=component["calculation_mode"], entered_value=component["entered_value"],
            normalized_per_kg=component["normalized_per_kg"], sort_order=index,
            metadata_json={"reference_qty_kg": original.reference_qty_kg}))
    sheet.version += 1
    db.commit(); db.refresh(version)
    return _cost_version_payload(version)


@router.post("/rm-costing/versions/{version_id}/activate")
def activate_cost_version(version_id: uuid.UUID, payload: VersionAction, db: Session = Depends(get_db),
                          plant_id: str = Depends(get_current_plant), current_user: dict = Depends(require_cost_administrator)):
    version = db.query(RmCostVersion).join(RmCostSheet).filter(RmCostVersion.id == version_id, RmCostSheet.plant_id == plant_id).with_for_update().first()
    if not version:
        raise HTTPException(status_code=404, detail="RM cost version not found")
    sheet = db.query(RmCostSheet).filter(RmCostSheet.id == version.sheet_id).with_for_update().first()
    if sheet.version != payload.expected_version or version.status not in {"DRAFT", "SCHEDULED"}:
        raise HTTPException(status_code=409, detail="Cost version changed; reload before activation")
    version.activation_reason = payload.reason.strip()
    if version.effective_from > date.today():
        version.status = "SCHEDULED"
    else:
        prior_effective = db.query(RmCostVersion).filter(
            RmCostVersion.sheet_id == sheet.id,
            RmCostVersion.id != version.id,
            RmCostVersion.status.in_(["ACTIVE", "SCHEDULED"]),
            RmCostVersion.effective_from <= version.effective_from,
        ).all()
        for previous in prior_effective:
            previous.status = "SUPERSEDED"; previous.effective_to = version.effective_from
        version.status = "ACTIVE"; sheet.current_version_id = version.id

    version.activated_by = _actor(current_user); version.activated_at = datetime.utcnow(); sheet.version += 1
    db.commit(); db.refresh(version)
    return _cost_version_payload(version)


class CostRestore(BaseModel):
    model_config = ConfigDict(extra="forbid")
    request_id: uuid.UUID
    expected_sheet_version: int = Field(gt=0)
    effective_from: date
    reason: str = Field(min_length=3, max_length=1000)


@router.post("/rm-costing/versions/{version_id}/restore-as-draft")
def restore_cost_as_draft(version_id: uuid.UUID, payload: CostRestore, db: Session = Depends(get_db),
                          plant_id: str = Depends(get_current_plant), current_user: dict = Depends(require_cost_administrator)):
    source = db.query(RmCostVersion).join(RmCostSheet).filter(RmCostVersion.id == version_id, RmCostSheet.plant_id == plant_id).first()
    if not source: raise HTTPException(status_code=404, detail="RM cost version not found")
    request_fingerprint = canonical_hash({"source_version_id": str(source.id), **payload.model_dump(mode="json", exclude={"request_id", "expected_sheet_version"})})
    sheet = db.query(RmCostSheet).filter(RmCostSheet.id == source.sheet_id).with_for_update().first()
    replay = db.query(RmCostVersion).filter(RmCostVersion.sheet_id == sheet.id, RmCostVersion.request_id == payload.request_id).first()
    if replay:
        if replay.request_fingerprint != request_fingerprint:
            raise HTTPException(status_code=409, detail="This restore request key already exists with different details")
        return _cost_version_payload(replay)
    if sheet.version != payload.expected_sheet_version:
        raise HTTPException(status_code=409, detail="Cost sheet changed; reload before restoring")
    latest = db.query(func.max(RmCostVersion.version_no)).filter(RmCostVersion.sheet_id == sheet.id).scalar() or 0
    restored = RmCostVersion(sheet_id=sheet.id, version_no=int(latest) + 1, request_id=payload.request_id,
        request_fingerprint=request_fingerprint,
        status="DRAFT", base_cost=source.base_cost, landed_cost=source.landed_cost, effective_from=payload.effective_from,
        change_reason=payload.reason, source=f"RESTORE_FROM_V{source.version_no}",
        content_hash=canonical_hash({"source_version_id": str(source.id), "effective_from": payload.effective_from, "reason": payload.reason}),
        created_by=_actor(current_user))
    db.add(restored); db.flush()
    for component in source.components or []:
        db.add(RmCostComponent(version_id=restored.id, component_type=component.component_type, label=component.label,
            calculation_mode=component.calculation_mode, entered_value=component.entered_value,
            normalized_per_kg=component.normalized_per_kg, sort_order=component.sort_order,
            metadata_json=dict(component.metadata_json or {})))
    sheet.version += 1; db.commit(); db.refresh(restored)
    return _cost_version_payload(restored)


class LabelJobCreate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    request_id: uuid.UUID
    lot_ids: list[uuid.UUID] = Field(min_length=1)
    copies: int = Field(default=1, gt=0, le=20)
    profile: str = Field(default="PAPER_LOT_4X2", pattern="^(PAPER_LOT_4X2|PAPER_LOT_A4)$")
    reprint_reason: Optional[str] = Field(default=None, max_length=1000)


@router.post("/label-jobs")
def create_label_job(payload: LabelJobCreate, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
                     current_user: dict = Depends(require_role(["Store", "PlantManager"]))):
    request_fingerprint = canonical_hash(payload.model_dump(mode="json", exclude={"request_id"}))
    replay = db.query(LabelPrintJob).filter(LabelPrintJob.plant_id == plant_id, LabelPrintJob.request_id == payload.request_id).first()
    if replay:
        if replay.request_fingerprint != request_fingerprint:
            raise HTTPException(status_code=409, detail="This label request key already exists with different details")
        return {"id": str(replay.id), "lot_ids": replay.lot_ids, "copies": replay.copies,
                "profile": replay.profile, "status": replay.status, "idempotent": True}
    records = db.query(LotLabelRecord).filter(LotLabelRecord.plant_id == plant_id, LotLabelRecord.reel_id.in_(payload.lot_ids)).all()
    if len(records) != len(set(payload.lot_ids)):
        raise HTTPException(status_code=404, detail="One or more saved lot labels were not found")
    prior = db.query(LabelPrintJob).filter(LabelPrintJob.plant_id == plant_id).all()
    previously_printed = any(set(row.lot_ids or []) & {str(value) for value in payload.lot_ids} for row in prior)
    if previously_printed and not payload.reprint_reason:
        raise HTTPException(status_code=422, detail="A reprint reason is required for previously printed lots")
    job = LabelPrintJob(plant_id=plant_id, request_id=payload.request_id, request_fingerprint=request_fingerprint, profile=payload.profile,
        lot_ids=[str(value) for value in payload.lot_ids], copies=payload.copies, status="GENERATED",
        reprint_reason=payload.reprint_reason, created_by=_actor(current_user))
    db.add(job); db.commit(); db.refresh(job)
    return {"id": str(job.id), "lot_ids": job.lot_ids, "copies": job.copies, "profile": job.profile,
        "status": job.status, "labels": [record.content_snapshot for record in records], "idempotent": False}


@router.get("/label-jobs/{job_id}/pdf")
def label_job_pdf(job_id: uuid.UUID, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    from ..services.procurement_documents import build_lot_label_pdf
    job = db.query(LabelPrintJob).filter_by(id=job_id, plant_id=plant_id).first()
    if not job:
        raise HTTPException(404, "Label job not found in this plant")
    records = {str(record.reel_id): record for record in db.query(LotLabelRecord).filter(LotLabelRecord.plant_id == plant_id, LotLabelRecord.reel_id.in_([uuid.UUID(value) for value in job.lot_ids])).all()}
    labels = []
    for lot_id in job.lot_ids:
        record = records.get(lot_id)
        if not record:
            raise HTTPException(409, "A saved label is unavailable; printing stopped")
        label = dict(record.content_snapshot)
        # Old snapshots predate width/form fields; preserve their QR and original weight.
        reel = db.get(PaperReel, uuid.UUID(lot_id))
        if reel:
            for key, value in {"width_mm": reel.width_mm, "physical_form": reel.physical_form, "source_reel_no": reel.source_reel_no}.items():
                if not label.get(key): label[key] = value
        labels.append(label)
    return StreamingResponse(BytesIO(build_lot_label_pdf(labels, job.copies, job.profile)), media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="lot-labels-{job.id}.pdf"'})


@router.get("/lots/{lot_id}")
def get_lot(lot_id: uuid.UUID, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    reel = db.query(PaperReel).filter(PaperReel.id == lot_id, PaperReel.plant_id == _plant_uuid(plant_id)).first()
    if not reel: raise HTTPException(status_code=404, detail="Paper lot not found")
    label = db.query(LotLabelRecord).filter(LotLabelRecord.reel_id == reel.id).first()
    receipt_line = db.query(PurchaseReceiptLine).filter(PurchaseReceiptLine.id == reel.purchase_receipt_line_id).first() if reel.purchase_receipt_line_id else None
    receipt = db.query(PurchaseReceipt).filter(PurchaseReceipt.id == receipt_line.receipt_id).first() if receipt_line else None
    return {"id": str(reel.id), "at_no": reel.reel_code, "source_reel_no": reel.source_reel_no,
        "physical_form": reel.physical_form, "original_received_kg": float(reel.inward_weight_kg),
        "current_kg": float(reel.current_weight_kg), "width_mm": reel.width_mm, "stock_status": reel.stock_status,
        "commercial_status": reel.commercial_status, "po_no": receipt.order.po_no if receipt and receipt.order else None,
        "receipt_id": str(receipt.id) if receipt else None, "grn_no": receipt.grn_no if receipt else None,
        "label": label.content_snapshot if label else None, "genealogy": reel.genealogy_metadata or {}}


def _csv_response(filename: str, rows: list[dict[str, Any]]) -> StreamingResponse:
    output = io.StringIO()
    columns = list(rows[0].keys()) if rows else ["message"]
    writer = csv.DictWriter(output, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    if rows:
        writer.writerows([{key: _export_cell(value) for key, value in row.items()} for row in rows])
    else:
        writer.writerow({"message": "No rows match the selected filter"})
    content = output.getvalue().encode("utf-8-sig")
    return StreamingResponse(iter([content]), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"', "X-Row-Count": str(len(rows))})


def _export_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, sort_keys=True, default=str)
    if isinstance(value, (uuid.UUID, Decimal)):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _xlsx_cell(value: Any) -> Any:
    """Preserve useful Excel types while preventing imported text formulas."""
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, sort_keys=True, default=str)
    if isinstance(value, uuid.UUID):
        value = str(value)
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _xlsx_response(filename: str, rows: list[dict[str, Any]]) -> StreamingResponse:
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = "Register"
    columns = list(rows[0].keys()) if rows else ["message"]
    sheet.append(columns)
    header_fill = PatternFill("solid", fgColor="0F172A")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    if rows:
        for row in rows:
            sheet.append([_xlsx_cell(row.get(column)) for column in columns])
    else:
        sheet.append(["No rows match the selected filter"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns, start=1):
        samples = [str(column), *[str(_export_cell(row.get(column))) for row in rows[:200]]]
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(max(map(len, samples)) + 2, 11), 42)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"', "X-Row-Count": str(len(rows))})


def _register_rows(register_name: str, db: Session, plant_id: str, current_user: dict) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if register_name == "po-lines":
        for order in db.query(PurchaseOrder).filter(PurchaseOrder.plant_id == plant_id).order_by(PurchaseOrder.created_at.desc()).all():
            for line in order.lines or []:
                rows.append({"plant_id": plant_id, "po_no": order.po_no, "category": order.category,
                    "revision": order.current_revision_no, "expected_date": order.expected_date, "po_date": (order.metadata_json or {}).get("po_date"),
                    "supplier": order.supplier_name_snapshot, "item_code": line.item.item_code if line.item else None,
                    "item_name": line.item.name if line.item else None, "ordered_qty": line.qty_ordered,
                    "received_qty": line.qty_received, "open_qty": float(line.qty_ordered)-float(line.qty_received)-float(line.qty_short_closed or 0),
                    "uom": line.uom, "expected_unit_count": line.expected_unit_count, "count_basis": line.count_basis,
                    "received_unit_count": line.received_unit_count,
                    "rate": line.unit_cost, "amount": float(line.qty_ordered)*float(line.unit_cost), "status": order.status,
                    "line_status": line.line_status, "specification": line.metadata_json, "document_terms": order.metadata_json, "created_by": order.created_by, "approved_by": order.approved_by})
    elif register_name == "receipts":
        for receipt in db.query(PurchaseReceipt).filter(PurchaseReceipt.plant_id == plant_id).order_by(PurchaseReceipt.created_at.desc()).all():
            payload = _receipt_payload(receipt, db)
            for line in payload["lines"]:
                rows.append({**{key: payload[key] for key in ("grn_no", "received_date", "receipt_kind", "po_no", "supplier_name", "invoice_no", "invoice_date", "commercial_status", "manual_reason")}, **line})
    elif register_name == "inward-lots":
        reels = db.query(PaperReel).filter(PaperReel.plant_id == _plant_uuid(plant_id)).order_by(PaperReel.created_at.desc()).all()
        for reel in reels:
            metadata = reel.inward_metadata or {}
            rows.append({"plant_id": plant_id, "inward_date": reel.inward_date, "at_no": reel.reel_code,
                "source_reel_no": reel.source_reel_no, "physical_form": reel.physical_form,
                "paper_id": reel.paper_id, "supplier": reel.supplier_name_snapshot, "po_no": metadata.get("po_no"),
                "po_revision": metadata.get("po_revision"), "grn_no": metadata.get("grn_no"),
                "invoice_no": metadata.get("invoice_no"), "width_mm": reel.width_mm,
                "original_kg": reel.inward_weight_kg, "current_kg": reel.current_weight_kg,
                "po_rate": metadata.get("po_rate"), "invoice_rate": metadata.get("invoice_rate"),
                "stock_status": reel.stock_status, "commercial_status": reel.commercial_status})
    elif register_name == "discrepancies":
        rows = [_discrepancy_payload(row, db) for row in db.query(PurchaseDiscrepancy).filter(PurchaseDiscrepancy.plant_id == plant_id).all()]
    elif register_name == "debit-notes":
        rows = [{key: value for key, value in _debit_payload(row).items() if key != "lines"}
                for row in db.query(PurchaseDebitNote).filter(PurchaseDebitNote.plant_id == plant_id).all()]
    elif register_name == "stock-policies":
        rows = [_policy_payload(row) for row in db.query(StockAlertPolicy).filter(StockAlertPolicy.plant_id == plant_id).all()]
    elif register_name == "rm-costing":
        costing = list_rm_costing(db=db, plant_id=plant_id, current_user=current_user)["items"]
        rows = costing
    elif register_name == "schedule":
        for plan in db.query(ProcurementPlan).filter(ProcurementPlan.plant_id == plant_id).all():
            for entry in plan.entries or []:
                rows.append({"plan": plan.name, "month": plan.month, "status": plan.status, "date": entry.entry_date,
                    "item_code": entry.item.item_code if entry.item else None, "item_name": entry.item.name if entry.item else None,
                    "supplier": entry.supplier_name_snapshot, "form": entry.material_form, "qty_kg": entry.qty_kg,
                    "expected_unit_count": entry.expected_unit_count, "converted_qty_kg": entry.converted_qty_kg})
    else:
        raise HTTPException(status_code=404, detail="Unknown procurement register")
    return rows


@router.get("/registers/{register_name}")
def view_register(register_name: str, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
                  current_user: dict = Depends(get_current_user)):
    rows = _register_rows(register_name, db, plant_id, current_user)
    return {"register": register_name, "row_count": len(rows),
            "columns": list(rows[0].keys()) if rows else [], "items": rows}


@router.get("/registers/{register_name}/export.csv")
def export_register_csv(register_name: str, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
                        current_user: dict = Depends(get_current_user)):
    rows = _register_rows(register_name, db, plant_id, current_user)
    return _csv_response(f"{register_name}-{plant_id}-{date.today().isoformat()}.csv", rows)


@router.get("/registers/{register_name}/export.xlsx")
def export_register_xlsx(register_name: str, db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant),
                         current_user: dict = Depends(get_current_user)):
    rows = _register_rows(register_name, db, plant_id, current_user)
    return _xlsx_response(f"{register_name}-{plant_id}-{date.today().isoformat()}.xlsx", rows)


@router.get("/registers/{register_name}/pdf")
def register_pdf(register_name: str, paper: str = Query(default="A3", pattern="^(A3|A4)$"), db: Session = Depends(get_db), plant_id: str = Depends(get_current_plant), current_user: dict = Depends(get_current_user)):
    from ..services.procurement_documents import build_register_pdf
    rows = _register_rows(register_name, db, plant_id, current_user)
    return StreamingResponse(BytesIO(build_register_pdf(register_name.replace("-", " ").title(), rows, paper)), media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{register_name}-{paper}.pdf"', "X-Row-Count": str(len(rows))})
